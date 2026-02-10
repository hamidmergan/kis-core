from io import BytesIO
from pathlib import Path
import json
import os
import re
import tempfile
from typing import Dict, List, Tuple
from urllib.parse import quote

import deepl
from docx import Document
from fpdf import FPDF
from dotenv import load_dotenv
from fastapi import Depends, FastAPI, File, HTTPException, Request, Security, UploadFile
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.openapi.docs import get_swagger_ui_html
from fastapi.security import APIKeyHeader
from pypdf import PdfReader
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

# Konfiguration aus .env laden (falls vorhanden)
load_dotenv()

# ---------------------------------------------------------------------------
# Globale API-Key-Absicherung
# ---------------------------------------------------------------------------
_OPEN_PATHS = {"/docs", "/redoc", "/openapi.json"}

api_key_header = APIKeyHeader(name="x-api-key", auto_error=False)


async def verify_api_key(
    request: Request,
    api_key: str = Security(api_key_header),
):
    """Prüft den API-Key für alle Endpunkte (außer Docs/OpenAPI)."""
    if request.url.path in _OPEN_PATHS:
        return
    token = os.getenv("KIS_API_TOKEN")
    if not token or api_key != token:
        raise HTTPException(status_code=401, detail="Ungültiger oder fehlender API-Key")


app = FastAPI(
    title="kis-thl - Dokumentübersetzung",
    dependencies=[Depends(verify_api_key)],
    docs_url=None,  # Standard-Docs deaktiviert → eigene Route unten
)


@app.get("/docs", include_in_schema=False)
async def custom_swagger_docs():
    """Swagger UI mit autocomplete=off & Passwort-Feld für den API-Key."""
    response = get_swagger_ui_html(
        openapi_url=app.openapi_url,
        title=f"{app.title} - Docs",
    )
    # Kleines Script injizieren: Auth-Inputs → type=password + autocomplete=off
    patch_script = """
<script>
(function () {
  const obs = new MutationObserver(() => {
    document.querySelectorAll('.auth-container input[type="text"]').forEach(el => {
      el.setAttribute('type', 'password');
      el.setAttribute('autocomplete', 'off');
    });
  });
  obs.observe(document.body, { childList: true, subtree: true });
})();
</script>"""
    html = response.body.decode()
    html = html.replace("</body>", patch_script + "\n</body>")
    return HTMLResponse(html)

TRANSLATION_PROVIDER = os.getenv("TRANSLATION_PROVIDER", "deepl")
DEEPL_API_KEY = os.getenv("DEEPL_API_KEY")
SOURCE_LANG = os.getenv("SOURCE_LANG", "DE")
TARGET_LANG_RAW = os.getenv("TARGET_LANG", "EN-US")
# DeepL benötigt "EN-US" oder "EN-GB", nicht "EN" - mappe "EN" zu "EN-US"
if TARGET_LANG_RAW.upper() == "EN":
    TARGET_LANG = "EN-US"
else:
    TARGET_LANG = TARGET_LANG_RAW

# DeepL Translator initialisieren (nur wenn Key vorhanden)
translator = None
if DEEPL_API_KEY:
    try:
        translator = deepl.Translator(DEEPL_API_KEY)
    except Exception as e:
        # Fehler beim Initialisieren wird später beim Übersetzen abgefangen
        pass


GLOSSARY_PATH = os.getenv("GLOSSARY_PATH", "glossary.json")
GLOSSARY: Dict[str, str] = {}
GLOSSARY_PATTERNS: List[Tuple[re.Pattern, str, str]] = []


def _normalize_glossary(data: Dict[str, str]) -> Dict[str, str]:
    if not isinstance(data, dict):
        return {}
    normalized: Dict[str, str] = {}
    for key, value in data.items():
        if isinstance(key, str) and isinstance(value, str):
            key = key.strip()
            if key:
                normalized[key] = value
    return normalized


def _load_glossary_from_file(path: str) -> Dict[str, str]:
    if not path:
        return {}
    try:
        with open(path, "r", encoding="utf-8") as handle:
            data = json.load(handle)
        return _normalize_glossary(data)
    except FileNotFoundError:
        return {}
    except Exception:
        return {}


def _write_glossary_to_file(path: str, glossary: Dict[str, str]) -> None:
    with open(path, "w", encoding="utf-8") as handle:
        json.dump(glossary, handle, ensure_ascii=False, indent=2)


def _build_glossary_patterns(glossary: Dict[str, str]) -> List[Tuple[re.Pattern, str, str]]:
    entries: List[Tuple[re.Pattern, str, str]] = []
    sorted_items = sorted(glossary.items(), key=lambda item: len(item[0]), reverse=True)
    for index, (source, target) in enumerate(sorted_items):
        if not source or not target:
            continue
        token = f"__GLOSSARY_{index}__"
        pattern = re.escape(source)
        if re.match(r"\w", source[0]):
            pattern = r"\b" + pattern
        if re.match(r"\w", source[-1]):
            pattern = pattern + r"\b"
        entries.append((re.compile(pattern), token, target))
    return entries


def _set_glossary(entries: Dict[str, str]) -> None:
    global GLOSSARY, GLOSSARY_PATTERNS
    GLOSSARY = entries
    GLOSSARY_PATTERNS = _build_glossary_patterns(entries)


def _apply_glossary_to_text(text: str) -> Tuple[str, Dict[str, str]]:
    if not GLOSSARY_PATTERNS:
        return text, {}
    updated = text
    replacements: Dict[str, str] = {}
    for pattern, token, target in GLOSSARY_PATTERNS:
        if pattern.search(updated):
            updated = pattern.sub(token, updated)
            replacements[token] = target
    return updated, replacements


def _restore_glossary(text: str, replacements: Dict[str, str]) -> str:
    restored = text
    for token, target in replacements.items():
        restored = restored.replace(token, target)
    return restored


def translate_text_with_glossary(text: str, source_lang: str, target_lang: str) -> str:
    if not translator:
        raise ValueError("Translator nicht initialisiert")
    if not text or not text.strip():
        return text
    prepared_text, replacements = _apply_glossary_to_text(text)
    translated = translator.translate_text(
        prepared_text, source_lang=source_lang, target_lang=target_lang
    ).text
    return _restore_glossary(translated, replacements)


def _split_markdown_by_code_fences(text: str) -> List[Tuple[bool, str]]:
    if not text:
        return [(False, text)]
    lines = text.splitlines(keepends=True)
    parts: List[Tuple[bool, str]] = []
    in_block = False
    fence_char = ""
    fence_len = 0
    buffer: List[str] = []

    for line in lines:
        if not in_block:
            match = re.match(r"^[ \t]{0,3}([`~]{3,})", line)
            if match:
                if buffer:
                    parts.append((False, "".join(buffer)))
                    buffer = []
                in_block = True
                fence_seq = match.group(1)
                fence_char = fence_seq[0]
                fence_len = len(fence_seq)
                buffer = [line]
            else:
                buffer.append(line)
        else:
            buffer.append(line)
            fence_pattern = r"^[ \t]{0,3}" + re.escape(fence_char) + "{" + str(fence_len) + r",}\s*$"
            if re.match(fence_pattern, line.rstrip("\n")):
                parts.append((True, "".join(buffer)))
                buffer = []
                in_block = False
                fence_char = ""
                fence_len = 0

    if buffer:
        parts.append((in_block, "".join(buffer)))

    return parts


def _translate_markdown_preserve_code_blocks(text: str, source_lang: str, target_lang: str) -> str:
    parts = _split_markdown_by_code_fences(text)
    translated_parts: List[str] = []
    for is_code, chunk in parts:
        if is_code or not chunk or not chunk.strip():
            translated_parts.append(chunk)
        else:
            translated_parts.append(
                translate_text_with_glossary(chunk, source_lang, target_lang)
            )
    return "".join(translated_parts)


_set_glossary(_load_glossary_from_file(GLOSSARY_PATH))


LATEX_MATH_ENVIRONMENTS = {
    "equation",
    "equation*",
    "align",
    "align*",
    "gather",
    "gather*",
    "multline",
    "multline*",
    "eqnarray",
    "eqnarray*",
    "math",
    "displaymath",
    "split",
    "flalign",
    "flalign*",
    "alignat",
    "alignat*",
}

LATEX_VERBATIM_ENVIRONMENTS = {
    "verbatim",
    "Verbatim",
    "verbatim*",
    "lstlisting",
    "minted",
}

LATEX_PROTECTED_ENVIRONMENTS = LATEX_MATH_ENVIRONMENTS | LATEX_VERBATIM_ENVIRONMENTS

LATEX_COMMAND_RULES = {
    "label": {"mandatory": ["protect"], "optional": "translate"},
    "ref": {"mandatory": ["protect"], "optional": "translate"},
    "eqref": {"mandatory": ["protect"], "optional": "translate"},
    "pageref": {"mandatory": ["protect"], "optional": "translate"},
    "autoref": {"mandatory": ["protect"], "optional": "translate"},
    "cref": {"mandatory": ["protect"], "optional": "translate"},
    "Cref": {"mandatory": ["protect"], "optional": "translate"},
    "cite": {"mandatory": ["protect"], "optional": "translate"},
    "citep": {"mandatory": ["protect"], "optional": "translate"},
    "citet": {"mandatory": ["protect"], "optional": "translate"},
    "citealt": {"mandatory": ["protect"], "optional": "translate"},
    "citealp": {"mandatory": ["protect"], "optional": "translate"},
    "citenum": {"mandatory": ["protect"], "optional": "translate"},
    "citeauthor": {"mandatory": ["protect"], "optional": "translate"},
    "citeyear": {"mandatory": ["protect"], "optional": "translate"},
    "citeyearpar": {"mandatory": ["protect"], "optional": "translate"},
    "url": {"mandatory": ["protect"], "optional": "protect"},
    "path": {"mandatory": ["protect"], "optional": "protect"},
    "href": {"mandatory": ["protect", "translate"], "optional": "translate"},
    "hyperref": {"mandatory": ["protect", "translate"], "optional": "translate"},
    "includegraphics": {"mandatory": ["protect"], "optional": "protect"},
    "input": {"mandatory": ["protect"], "optional": "protect"},
    "include": {"mandatory": ["protect"], "optional": "protect"},
    "includeonly": {"mandatory": ["protect"], "optional": "protect"},
    "bibliography": {"mandatory": ["protect"], "optional": "protect"},
    "bibliographystyle": {"mandatory": ["protect"], "optional": "protect"},
    "addbibresource": {"mandatory": ["protect"], "optional": "protect"},
    "documentclass": {"mandatory": ["protect"], "optional": "protect"},
    "usepackage": {"mandatory": ["protect"], "optional": "protect"},
    "lstinputlisting": {"mandatory": ["protect"], "optional": "protect"},
    "verbatiminput": {"mandatory": ["protect"], "optional": "protect"},
    "graphicspath": {"mandatory": ["protect"], "optional": "protect"},
}

LATEX_VERBATIM_COMMANDS = {"verb", "lstinline"}


def _latex_is_escaped(text: str, idx: int) -> bool:
    backslashes = 0
    j = idx - 1
    while j >= 0 and text[j] == "\\":
        backslashes += 1
        j -= 1
    return (backslashes % 2) == 1


def _latex_read_group(text: str, start_idx: int, open_char: str, close_char: str):
    if start_idx >= len(text) or text[start_idx] != open_char:
        return None, start_idx
    depth = 0
    i = start_idx
    while i < len(text):
        ch = text[i]
        if ch == open_char and not _latex_is_escaped(text, i):
            depth += 1
        elif ch == close_char and not _latex_is_escaped(text, i):
            depth -= 1
            if depth == 0:
                group = text[start_idx : i + 1]
                return group, i + 1
        i += 1
    return None, start_idx


def _latex_consume_whitespace(text: str, idx: int):
    start = idx
    while idx < len(text) and text[idx].isspace():
        idx += 1
    return text[start:idx], idx


def _latex_find_math_end(text: str, start_idx: int, delimiter: str):
    i = start_idx
    delim_len = len(delimiter)
    while i < len(text):
        if text.startswith(delimiter, i) and not _latex_is_escaped(text, i):
            return i + delim_len
        i += 1
    return len(text)


def _latex_find_environment_end(text: str, start_idx: int, env_name: str):
    end_tag = f"\\end{{{env_name}}}"
    pos = text.find(end_tag, start_idx)
    if pos == -1:
        return len(text)
    return pos + len(end_tag)


def _translate_latex_preserve_structure(text: str, source_lang: str, target_lang: str) -> str:
    if not text or not text.strip():
        return text

    out: List[str] = []
    buffer: List[str] = []

    def flush_buffer():
        if not buffer:
            return
        chunk = "".join(buffer)
        if chunk.strip():
            out.append(_translate_text_preserve_linebreaks(chunk, source_lang, target_lang))
        else:
            out.append(chunk)
        buffer.clear()

    i = 0
    while i < len(text):
        ch = text[i]

        # Comments
        if ch == "%" and not _latex_is_escaped(text, i):
            flush_buffer()
            end = text.find("\n", i)
            if end == -1:
                out.append(text[i:])
                return "".join(out)
            out.append(text[i : end + 1])
            i = end + 1
            continue

        # Inline/display math with $...$ or $$...$$
        if ch == "$" and not _latex_is_escaped(text, i):
            flush_buffer()
            if i + 1 < len(text) and text[i + 1] == "$":
                end = _latex_find_math_end(text, i + 2, "$$")
                out.append(text[i:end])
                i = end
            else:
                end = _latex_find_math_end(text, i + 1, "$")
                out.append(text[i:end])
                i = end
            continue

        # Math delimiters \(...\) and \[...\]
        if text.startswith("\\(", i) or text.startswith("\\[", i):
            flush_buffer()
            closing = "\\)" if text.startswith("\\(", i) else "\\]"
            end = text.find(closing, i + 2)
            if end == -1:
                out.append(text[i:])
                return "".join(out)
            out.append(text[i : end + 2])
            i = end + 2
            continue

        # Braces are kept as structural tokens
        if ch in "{}":
            flush_buffer()
            out.append(ch)
            i += 1
            continue

        # Commands
        if ch == "\\":
            flush_buffer()
            cmd_start = i
            i += 1
            if i >= len(text):
                out.append("\\")
                break
            if text[i].isalpha():
                j = i
                while j < len(text) and text[j].isalpha():
                    j += 1
                cmd_name = text[i:j]
                cmd_token = text[cmd_start:j]

                # Verbatim-like inline commands
                if cmd_name in LATEX_VERBATIM_COMMANDS:
                    k = j
                    if k < len(text) and text[k] == "*":
                        k += 1
                        cmd_token = text[cmd_start:k]
                    if k >= len(text):
                        out.append(cmd_token)
                        i = k
                        continue
                    delim = text[k]
                    end = text.find(delim, k + 1)
                    if end == -1:
                        out.append(text[cmd_start:])
                        return "".join(out)
                    out.append(text[cmd_start : end + 1])
                    i = end + 1
                    continue

                # \begin{env} / \end{env}
                if cmd_name in ("begin", "end"):
                    out.append(cmd_token)
                    ws, pos = _latex_consume_whitespace(text, j)
                    if ws:
                        out.append(ws)
                    group, end_pos = _latex_read_group(text, pos, "{", "}")
                    if not group:
                        i = j
                        continue
                    out.append(group)
                    i = end_pos
                    if cmd_name == "begin":
                        env_name = group[1:-1]
                        if env_name in LATEX_PROTECTED_ENVIRONMENTS:
                            end_env = _latex_find_environment_end(text, i, env_name)
                            out.append(text[i:end_env])
                            i = end_env
                    continue

                # Other commands with argument rules
                rule = LATEX_COMMAND_RULES.get(cmd_name)
                out.append(cmd_token)
                i = j
                if rule:
                    # Optional args
                    while True:
                        ws, i = _latex_consume_whitespace(text, i)
                        if ws:
                            out.append(ws)
                        if i < len(text) and text[i] == "[":
                            group, end_pos = _latex_read_group(text, i, "[", "]")
                            if not group:
                                break
                            if rule.get("optional") == "protect":
                                out.append(group)
                            else:
                                inner = group[1:-1]
                                translated = _translate_latex_preserve_structure(
                                    inner, source_lang, target_lang
                                )
                                out.append("[" + translated + "]")
                            i = end_pos
                            continue
                        break

                    # Mandatory args
                    arg_modes = rule.get("mandatory", [])
                    for mode in arg_modes:
                        ws, i = _latex_consume_whitespace(text, i)
                        if ws:
                            out.append(ws)
                        if i < len(text) and text[i] == "{":
                            group, end_pos = _latex_read_group(text, i, "{", "}")
                            if not group:
                                break
                            if mode == "protect":
                                out.append(group)
                            else:
                                inner = group[1:-1]
                                translated = _translate_latex_preserve_structure(
                                    inner, source_lang, target_lang
                                )
                                out.append("{" + translated + "}")
                            i = end_pos
                        else:
                            break
                continue

            # Single-character command like \%, \_, \{, \}
            out.append(text[cmd_start : i + 1])
            i += 1
            continue

        buffer.append(ch)
        i += 1

    flush_buffer()
    return "".join(out)


ALLOWED_EXTENSIONS = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".md": "text/markdown",
    ".tex": "application/x-tex",
}

DOC_TRANSLATION_EXTENSIONS = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".pdf": "application/pdf",
}

MAX_UPLOAD_BYTES = 200 * 1024 * 1024  # 200 MB Limit


def _iter_docx_paragraphs(container):
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_docx_paragraphs(cell)


def _iter_docx_paragraphs_with_header_footer(doc):
    yield from _iter_docx_paragraphs(doc)
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            yield from _iter_docx_paragraphs(header)
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            yield from _iter_docx_paragraphs(footer)


def translate_docx(content: bytes, source_lang: str, target_lang: str) -> bytes:
    """
    Übersetzt Text in einer DOCX-Datei.
    Layout, Formatierung und Struktur bleiben erhalten.
    Strategie: Absatz- und Run-basiert, um Formatierungen zu bewahren.
    Inklusive Header und Footer.
    """
    if not translator:
        raise ValueError("Translator nicht initialisiert")
    
    try:
        # DOCX aus Bytes öffnen
        doc = Document(BytesIO(content))
        
        # Absätze einzeln übersetzen, um Runs/Formatierung (z.B. Bold) zu erhalten.
        any_text = False
        seen_paragraphs = set()
        for para in _iter_docx_paragraphs_with_header_footer(doc):
            if id(para) in seen_paragraphs:
                continue
            seen_paragraphs.add(id(para))
            text = para.text
            if not text or not text.strip():
                continue
            any_text = True
            try:
                _translate_paragraph_preserve_runs(
                    para, source_lang, target_lang
                )
            except Exception as e:
                raise ValueError(
                    f"Fehler beim Übersetzen eines Absatzes: {str(e)}"
                )

        if not any_text:
            # Keine Text-Absätze gefunden, Datei unverändert zurückgeben
            output = BytesIO()
            doc.save(output)
            return output.getvalue()
        
        # Zurück zu Bytes konvertieren
        output = BytesIO()
        doc.save(output)
        output.seek(0)
        return output.getvalue()
        
    except ValueError:
        raise
    except Exception as e:
        import traceback
        raise ValueError(f"Fehler beim Verarbeiten der DOCX-Datei: {str(e)}\n{traceback.format_exc()}")


def _iter_pptx_paragraphs(shape):
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            yield paragraph
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    yield paragraph
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for subshape in shape.shapes:
            yield from _iter_pptx_paragraphs(subshape)


def _translate_run_text_preserve_linebreaks(text: str, source_lang: str, target_lang: str) -> str:
    """Übersetzt Text und erhält Zeilenumbrüche, führende/nachfolgende Leerzeichen und Tabs."""
    # Tabs durch Platzhalter ersetzen (DeepL würde sie sonst entfernen)
    TAB_PLACEHOLDER = "⟨TAB⟩"
    text_with_placeholder = text.replace("\t", TAB_PLACEHOLDER)
    
    lines = text_with_placeholder.split("\n")
    translated_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped:
            # Führende und nachfolgende Whitespace (Leerzeichen, Tabs) erfassen
            leading_ws = line[:len(line) - len(line.lstrip())]
            trailing_ws = line[len(line.rstrip()):]
            # Nur den gestrippten Text übersetzen
            translated = translate_text_with_glossary(
                stripped, source_lang, target_lang
            )
            # Whitespace wieder hinzufügen
            translated_lines.append(leading_ws + translated + trailing_ws)
        else:
            # Leere Zeile oder nur Whitespace: unverändert lassen
            translated_lines.append(line)
    
    result = "\n".join(translated_lines)
    # Platzhalter wieder durch Tabs ersetzen
    return result.replace(TAB_PLACEHOLDER, "\t")


def _translate_paragraph_preserve_runs(paragraph, source_lang: str, target_lang: str):
    runs = list(paragraph.runs)
    if not runs:
        paragraph.text = _translate_text_preserve_linebreaks(
            paragraph.text, source_lang, target_lang
        )
        return

    if len(runs) == 1:
        runs[0].text = _translate_text_preserve_linebreaks(
            runs[0].text, source_lang, target_lang
        )
        return

    # Mehrere Runs: run-weise übersetzen, um Schriftgrößen/Formatierung zu erhalten
    for run in runs:
        if run.text.strip():
            run.text = _translate_run_text_preserve_linebreaks(
                run.text, source_lang, target_lang
            )


def _translate_text_preserve_linebreaks(text: str, source_lang: str, target_lang: str) -> str:
    """Übersetzt Text und erhält Zeilenumbrüche, führende/nachfolgende Leerzeichen und Tabs."""
    # Tabs durch Platzhalter ersetzen (DeepL würde sie sonst entfernen)
    TAB_PLACEHOLDER = "⟨TAB⟩"
    text_with_placeholder = text.replace("\t", TAB_PLACEHOLDER)
    
    lines = text_with_placeholder.split("\n")
    translated_lines = []
    for line in lines:
        stripped = line.strip()
        if stripped:
            # Führende und nachfolgende Whitespace (Leerzeichen, Tabs) erfassen
            leading_ws = line[:len(line) - len(line.lstrip())]
            trailing_ws = line[len(line.rstrip()):]
            # Nur den gestrippten Text übersetzen
            translated = translate_text_with_glossary(
                stripped, source_lang, target_lang
            )
            # Whitespace wieder hinzufügen
            translated_lines.append(leading_ws + translated + trailing_ws)
        else:
            # Leere Zeile oder nur Whitespace: unverändert lassen
            translated_lines.append(line)
    
    result = "\n".join(translated_lines)
    # Platzhalter wieder durch Tabs ersetzen
    return result.replace(TAB_PLACEHOLDER, "\t")


def translate_pptx(content: bytes, source_lang: str, target_lang: str) -> bytes:
    """
    Übersetzt Text in einer PPTX-Datei.
    Text wird absatzweise in Blöcken übersetzt, Layout bleibt weitgehend erhalten.
    Inkl. Footer, Notizen und Slide-Master.
    """
    if not translator:
        raise ValueError("Translator nicht initialisiert")

    try:
        presentation = Presentation(BytesIO(content))

        any_text = False
        
        # 1. Normale Folien-Shapes
        for slide in presentation.slides:
            for shape in slide.shapes:
                for para in _iter_pptx_paragraphs(shape):
                    raw_text = para.text
                    if not raw_text or not raw_text.strip():
                        continue
                    any_text = True
                    try:
                        _translate_paragraph_preserve_runs(
                            para, source_lang, target_lang
                        )
                    except Exception as e:
                        raise ValueError(f"Fehler beim Übersetzen von PPTX-Absatz: {str(e)}")
            
            # 2. Notizen-Folien (Speaker Notes)
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
                for shape in notes_slide.shapes:
                    for para in _iter_pptx_paragraphs(shape):
                        raw_text = para.text
                        if not raw_text or not raw_text.strip():
                            continue
                        any_text = True
                        try:
                            _translate_paragraph_preserve_runs(
                                para, source_lang, target_lang
                            )
                        except Exception:
                            pass  # Notizen-Fehler ignorieren
        
        # 3. Slide-Master (Footer, Header, etc.)
        for slide_master in presentation.slide_masters:
            for shape in slide_master.shapes:
                for para in _iter_pptx_paragraphs(shape):
                    raw_text = para.text
                    if not raw_text or not raw_text.strip():
                        continue
                    any_text = True
                    try:
                        _translate_paragraph_preserve_runs(
                            para, source_lang, target_lang
                        )
                    except Exception:
                        pass  # Master-Fehler ignorieren
            
            # Slide-Layouts im Master
            for layout in slide_master.slide_layouts:
                for shape in layout.shapes:
                    for para in _iter_pptx_paragraphs(shape):
                        raw_text = para.text
                        if not raw_text or not raw_text.strip():
                            continue
                        any_text = True
                        try:
                            _translate_paragraph_preserve_runs(
                                para, source_lang, target_lang
                            )
                        except Exception:
                            pass  # Layout-Fehler ignorieren

        if not any_text:
            output = BytesIO()
            presentation.save(output)
            return output.getvalue()

        output = BytesIO()
        presentation.save(output)
        output.seek(0)
        return output.getvalue()
    except ValueError:
        raise
    except Exception as e:
        import traceback
        raise ValueError(f"Fehler beim Verarbeiten der PPTX-Datei: {str(e)}\n{traceback.format_exc()}")


def translate_pdf(content: bytes, source_lang: str, target_lang: str) -> bytes:
    """
    Übersetzt extrahierten Text aus einer PDF-Datei und baut ein einfaches PDF neu auf.
    Layout kann abweichen, Text geht aber nicht verloren. Ist kein Text extrahierbar,
    wird die Originaldatei zurückgegeben.
    """
    if not translator:
        raise ValueError("Translator nicht initialisiert")

    try:
        reader = PdfReader(BytesIO(content))
        translated_pages = []
        any_text = False

        for page in reader.pages:
            text = page.extract_text() or ""
            stripped = text.strip()
            if stripped:
                any_text = True
                translated = translate_text_with_glossary(
                    stripped, source_lang, target_lang
                )
            else:
                translated = ""
            translated_pages.append(translated)

        if not any_text:
            # Kein extrahierbarer Text: Datei unverändert zurückgeben
            return content

        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.set_font("Arial", size=12)

        for page_text in translated_pages:
            pdf.add_page()
            if page_text:
                safe_text = page_text.encode("latin-1", errors="replace").decode("latin-1")
                pdf.multi_cell(0, 10, txt=safe_text)

        raw_pdf = pdf.output(dest="S")
        if isinstance(raw_pdf, str):
            return raw_pdf.encode("latin-1")
        return bytes(raw_pdf)

    except ValueError:
        raise
    except Exception as e:
        import traceback
        raise ValueError(f"Fehler beim Verarbeiten der PDF-Datei: {str(e)}\n{traceback.format_exc()}")


def translate_document_via_deepl(content: bytes, filename: str, source_lang: str, target_lang: str) -> bytes:
    """
    Nutzt DeepL Document Translation API direkt (unterstützt PDF/DOCX/PPTX).
    """
    if not translator:
        raise ValueError("DeepL Translator nicht initialisiert")

    try:
        # Neuere deepl SDKs haben translate_document_from_bytes, ältere nicht.
        if hasattr(translator, "translate_document_from_bytes"):
            handle = translator.translate_document_from_bytes(
                content,
                filename,
                target_lang=target_lang,
                source_lang=source_lang,
            )
        else:
            # Fallback: temporäre Datei schreiben und translate_document_from_filepath nutzen.
            suffix = Path(filename).suffix or ""
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                tmp.write(content)
                tmp_path = tmp.name
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp_out:
                tmp_out_path = tmp_out.name
            try:
                result = translator.translate_document_from_filepath(
                    tmp_path,
                    output_path=tmp_out_path,
                    target_lang=target_lang,
                    source_lang=source_lang,
                )
                if hasattr(result, "document_id"):
                    translator.translate_document_wait(result)
                    translated_path = translator.translate_document_download(result, tmp_out_path)
                    if isinstance(translated_path, (bytes, bytearray)):
                        return bytes(translated_path)
                    tmp_out_path = translated_path if translated_path else tmp_out_path
                # Wenn kein Handle, ist Ergebnis synchron und liegt unter tmp_out_path
                with open(tmp_out_path, "rb") as f:
                    return f.read()
            finally:
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass
                try:
                    os.remove(tmp_out_path)
                except OSError:
                    pass

        translator.translate_document_wait(handle)
        result = translator.translate_document_download(handle)
        return result if isinstance(result, (bytes, bytearray)) else bytes(result)
    except Exception as e:
        import traceback
        raise ValueError(f"DeepL Document Translation Fehler: {str(e)}\n{traceback.format_exc()}")


def translate_filename(filename: str, source_lang: str, target_lang: str) -> str:
    """Übersetzt den Dateinamen (ohne Erweiterung)."""
    if not translator:
        return filename
    
    stem = Path(filename).stem
    ext = Path(filename).suffix
    
    # Nur übersetzen, wenn es Text gibt
    if not stem.strip():
        return filename
    
    try:
        translated = translate_text_with_glossary(
            stem, source_lang, target_lang
        )
        # Leerzeichen durch Unterstriche ersetzen (für Dateinamen)
        translated = translated.replace(" ", "_")
        return f"{translated}{ext}"
    except Exception:
        # Bei Fehler: Originalname behalten
        return filename


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/glossary")
def get_glossary():
    return {"count": len(GLOSSARY), "path": GLOSSARY_PATH, "entries": GLOSSARY}


@app.post("/glossary")
def set_glossary(entries: Dict[str, str], persist: bool = False):
    normalized = _normalize_glossary(entries)
    _set_glossary(normalized)
    if persist:
        if not GLOSSARY_PATH:
            raise HTTPException(status_code=400, detail="GLOSSARY_PATH nicht gesetzt")
        try:
            _write_glossary_to_file(GLOSSARY_PATH, normalized)
        except Exception as e:
            raise HTTPException(
                status_code=500,
                detail=f"Glossar konnte nicht gespeichert werden: {str(e)}",
            )
    return {"count": len(GLOSSARY), "path": GLOSSARY_PATH}


@app.post("/glossary/reload")
def reload_glossary():
    _set_glossary(_load_glossary_from_file(GLOSSARY_PATH))
    return {"count": len(GLOSSARY), "path": GLOSSARY_PATH}


@app.post("/translate")
async def translate(request: Request, file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Dateiname fehlt")

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Nur .docx, .pptx, .md oder .tex erlaubt")

    content_length = request.headers.get("content-length")
    if content_length is not None and int(content_length) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Datei zu groß (max 200 MB)")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Datei ist leer")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Datei zu groß (max 200 MB)")

    # TARGET_LANG sollte bereits beim Laden gemappt sein (EN -> EN-US)
    # Zur absoluten Sicherheit nochmal prüfen und mappen
    target_lang_final = TARGET_LANG.strip()
    if target_lang_final.upper() in ("EN", ""):
        target_lang_final = "EN-US"
    
    if not translator:
        raise HTTPException(
            status_code=500,
            detail="DeepL API Key fehlt oder Translator konnte nicht initialisiert werden",
        )
    
    # Dateinamen übersetzen (einmal für alle Fälle)
    translated_filename = translate_filename(file.filename, SOURCE_LANG, target_lang_final)
    download_name = f"{Path(translated_filename).stem}_translated{ext}"
    
    # Markdown-Dateien übersetzen
    if ext == ".md":
        try:
            # Markdown als UTF-8 Text dekodieren
            text = content.decode("utf-8")
            translated_text = _translate_markdown_preserve_code_blocks(
                text, SOURCE_LANG, target_lang_final
            )
            # Zurueck zu Bytes encoden
            content = translated_text.encode("utf-8")
        except Exception as e:
            raise HTTPException(
                status_code=500, detail=f"uebersetzungsfehler: {str(e)}"
            )
    elif ext == ".tex":
        try:
            text = content.decode("utf-8")
            translated_text = _translate_latex_preserve_structure(
                text, SOURCE_LANG, target_lang_final
            )
            content = translated_text.encode("utf-8")
        except Exception as e:
            raise HTTPException(
                status_code=500, detail=f"LaTeX-Übersetzungsfehler: {str(e)}"
            )
    elif ext == ".docx":
        try:
            content = translate_docx(content, SOURCE_LANG, target_lang_final)
        except ValueError as e:
            raise HTTPException(
                status_code=500, detail=f"DOCX-Übersetzungsfehler: {str(e)}"
            )
        except Exception as e:
            # Detaillierte Fehlermeldung für Debugging
            import traceback
            error_msg = str(e)
            error_trace = traceback.format_exc()
            # Fehler auch in Console loggen
            print(f"DOCX-Fehler: {error_msg}")
            print(f"Traceback: {error_trace}")
            raise HTTPException(
                status_code=500, detail=f"DOCX-Verarbeitungsfehler: {error_msg}"
            )
    # PPTX-Dateien übersetzen
    elif ext == ".pptx":
        try:
            content = translate_pptx(content, SOURCE_LANG, target_lang_final)
        except ValueError as e:
            raise HTTPException(
                status_code=500, detail=f"PPTX-Übersetzungsfehler: {str(e)}"
            )
        except Exception as e:
            import traceback
            error_msg = str(e)
            error_trace = traceback.format_exc()
            print(f"PPTX-Fehler: {error_msg}")
            print(f"Traceback: {error_trace}")
            raise HTTPException(
                status_code=500, detail=f"PPTX-Verarbeitungsfehler: {error_msg}"
            )
    else:
        # Fallback: sollte wegen ALLOWED_EXTENSIONS nicht erreicht werden
        raise HTTPException(status_code=400, detail="Dateityp wird nicht unterstützt")

    encoded_filename = quote(download_name, safe='')
    return StreamingResponse(
        iter([content]),
        media_type=ALLOWED_EXTENSIONS[ext],
        headers={
            "Content-Disposition": f'attachment; filename="{encoded_filename}"; filename*=UTF-8\'\'{encoded_filename}'
        },
    )


@app.post("/translate-document")
async def translate_document(request: Request, file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="Dateiname fehlt")

    ext = Path(file.filename).suffix.lower()
    if ext not in DOC_TRANSLATION_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Nur .docx, .pptx oder .pdf erlaubt")

    content_length = request.headers.get("content-length")
    if content_length is not None and int(content_length) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Datei zu groß (max 200 MB)")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Datei ist leer")
    if len(content) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="Datei zu groß (max 200 MB)")

    target_lang_final = TARGET_LANG.strip()
    if target_lang_final.upper() in ("EN", ""):
        target_lang_final = "EN-US"

    try:
        translated_bytes = translate_document_via_deepl(content, file.filename, SOURCE_LANG, target_lang_final)
        # Dateinamen übersetzen
        translated_filename = translate_filename(file.filename, SOURCE_LANG, target_lang_final)
        download_name = f"{Path(translated_filename).stem}_translated{ext}"
    except ValueError as e:
        raise HTTPException(status_code=500, detail=f"Document-Übersetzungsfehler: {str(e)}")
    except Exception as e:
        import traceback
        error_msg = str(e)
        error_trace = traceback.format_exc()
        print(f"Document-Fehler: {error_msg}")
        print(f"Traceback: {error_trace}")
        raise HTTPException(status_code=500, detail=f"Document-Verarbeitungsfehler: {error_msg}")

    encoded_filename = quote(download_name, safe='')
    return StreamingResponse(
        iter([translated_bytes]),
        media_type=DOC_TRANSLATION_EXTENSIONS[ext],
        headers={
            "Content-Disposition": f'attachment; filename="{encoded_filename}"; filename*=UTF-8\'\'{encoded_filename}'
        },
    )
