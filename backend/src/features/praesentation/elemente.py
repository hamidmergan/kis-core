"""Tables and charts in the template's colour, sized to whatever count of rows, columns and series they hold."""

from __future__ import annotations

from loguru import logger
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from src.features.praesentation import farben as fb
from src.features.praesentation.farben import Farben
from src.features.praesentation.schemas import MAX_TEXT_ITEMS, Body
from src.features.praesentation.vorlage import EMU_PER_PT, Feld
from src.services.screencast.pptx import _load_font

_CHART = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}
# PowerPoint's own "No Style, No Grid", so the fills set here are the only ones drawn.
_OHNE_STIL = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
ZEILE = 1.2
# A line that fills its box exactly wraps in PowerPoint, so text is laid out on this share of its box.
_SPIEL = 0.95
# Below this even a table of eight columns is no longer readable from the back of a room.
MIN_PT = 10.0
# A template with large body text would blow a small table up to poster size.
MAX_PT = 20.0
# Chart text a notch below the body text, never tiny and never louder than the heading.
_DIAGRAMM_ANTEIL = 0.85
_DIAGRAMM_PT = (10.0, 18.0)
# Value labels crowd a chart beyond this many points, the axis then reads better.
_MAX_BESCHRIFTETE_PUNKTE = 12
# A cell's air in multiples of the text size: at least the snug one, the roomy one where the table has space.
_SEITE = (0.5, 0.9)
_OBEN = (0.3, 0.75)


def zeilen_hart(text: str, font, breite_pt: float) -> int:
    """Greedy word wrap that also breaks a word too long for one line, as PowerPoint does."""
    if breite_pt <= 0:
        return 1
    n, aktuell = 1, 0.0
    leer = font.getlength(" ")
    for wort in text.split():
        w = font.getlength(wort)
        if aktuell and aktuell + leer + w > breite_pt:
            n += 1
            aktuell = 0.0
        if w > breite_pt:
            n += int(w // breite_pt)
            w = w % breite_pt
        aktuell += (leer if aktuell else 0.0) + w
    return n


def schrift(typeface: str, groesse: float, fett: bool = False):
    name = f"{typeface}:bold" if fett else typeface
    return _load_font(name, round(groesse))


def setze_lauf(lauf, groesse: float, farbe: RGBColor | None = None, fett: bool = False, typeface: str = "") -> None:
    lauf.font.size = Pt(groesse)
    lauf.font.bold = fett
    if farbe is not None:
        lauf.font.color.rgb = farbe
    if typeface:
        lauf.font.name = typeface


def tabelle(slide, feld: Feld, body: Body, f: Farben, typeface: str) -> None:
    """Header in full colour, rows banded at five percent, widths and size fitted to the content."""
    kopf = body.header[:MAX_TEXT_ITEMS]
    spalten = min(MAX_TEXT_ITEMS, max([len(kopf)] + [len(r) for r in body.rows]))
    zellen = ([kopf + [""] * (spalten - len(kopf))] if kopf else []) + [
        (r + [""] * spalten)[:spalten] for r in body.rows]
    groesse, breiten, hoehen, seite, oben = _tabellenmass(zellen, feld, typeface or "Calibri", bool(kopf))
    rahmen = slide.shapes.add_table(len(zellen), spalten, feld.left, feld.top, feld.width, sum(hoehen))
    t = rahmen.table
    t._tbl.tblPr.find(qn("a:tableStyleId")).text = _OHNE_STIL
    t.first_row = bool(kopf)
    t.horz_banding = False
    for c, b in enumerate(breiten):
        t.columns[c].width = Emu(b)
    for r, h in enumerate(hoehen):
        t.rows[r].height = Emu(h)
    for r, zeile in enumerate(zellen):
        ist_kopf = bool(kopf) and r == 0
        band = r - (1 if kopf else 0)
        for c, text in enumerate(zeile):
            zelle = t.cell(r, c)
            zelle.fill.solid()
            if ist_kopf:
                zelle.fill.fore_color.rgb = f.ton(fb.VOLL)
            elif band % 2 == 1:
                zelle.fill.fore_color.rgb = f.ton(fb.STREIFEN)
            else:
                zelle.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            zelle.vertical_anchor = MSO_ANCHOR.MIDDLE
            zelle.margin_left = zelle.margin_right = Emu(int(groesse * seite * EMU_PER_PT))
            zelle.margin_top = zelle.margin_bottom = Emu(int(groesse * oben * EMU_PER_PT))
            zelle.text = text
            for lauf in zelle.text_frame.paragraphs[0].runs:
                setze_lauf(lauf, groesse, f.schrift_auf() if ist_kopf else RGBColor(0, 0, 0), ist_kopf, typeface)


def _spaltenbreiten(zellen: list[list[str]], gesamt_pt: float, font, rand_pt: float) -> list[float] | None:
    """Widths by text length when all fits one line, else each column its longest word first, None when even that overflows."""
    n = len(zellen[0])
    wort = [max((font.getlength(w) for z in zellen for w in z[c].split()), default=0.0) / _SPIEL + rand_pt
            for c in range(n)]
    rest = gesamt_pt - sum(wort)
    if rest < 0:
        return None
    lang = [max(font.getlength(z[c]) for z in zellen) / _SPIEL + rand_pt for c in range(n)]
    if sum(lang) <= gesamt_pt:
        # Every cell fits one line, so the columns simply share the width by their length.
        return [gesamt_pt * l / sum(lang) for l in lang]
    # Long text may wrap, so a column claims at most twice the mean however long its sentence.
    deckel = 2 * sum(lang) / n
    bedarf = [max(0.0, min(l, deckel) - w) + 1 for l, w in zip(lang, wort)]
    return [w + rest * b / sum(bedarf) for w, b in zip(wort, bedarf)]


def _tabellenmass(zellen: list[list[str]], feld: Feld, typeface: str,
                  kopf: bool) -> tuple[float, list[int], list[int], float, float]:
    """The largest size at which every row fits the field and no word breaks, with widths, heights and cell air.

    The roomy side air is taken when it breaks no extra line, the air above and below grows with the
    room the rows leave free.
    """
    breite_pt, hoehe_pt = feld.width / EMU_PER_PT, feld.height / EMU_PER_PT
    groesse = max(MIN_PT, float(min(MAX_PT, round(feld.size_pt or 16))))
    while True:
        # Measured in bold, so the header row fits as well as the body.
        font = schrift(typeface, groesse, fett=kopf)
        eng = _spaltenbreiten(zellen, breite_pt, font, 2 * _SEITE[0] * groesse)
        if eng is None and groesse > MIN_PT:
            groesse -= 1
            continue
        eng = eng or [breite_pt / len(zellen[0])] * len(zellen[0])
        zeilen = _zeilenzahlen(zellen, eng, groesse, _SEITE[0], typeface, kopf)
        breiten, seite = eng, _SEITE[0]
        weit = _spaltenbreiten(zellen, breite_pt, font, 2 * _SEITE[1] * groesse)
        if weit is not None and _zeilenzahlen(zellen, weit, groesse, _SEITE[1], typeface, kopf) == zeilen:
            breiten, seite = weit, _SEITE[1]
        text = sum(zeilen) * groesse * ZEILE
        if text + 2 * _OBEN[0] * groesse * len(zellen) <= hoehe_pt or groesse <= MIN_PT:
            frei = (hoehe_pt - text) / (2 * groesse * len(zellen))
            oben = max(_OBEN[0], min(_OBEN[1], frei))
            if frei < _OBEN[0]:
                logger.info("Tabelle ragt auch bei {} pt über ihr Feld", MIN_PT)
            emu = [int(b * EMU_PER_PT) for b in breiten]
            emu[-1] += feld.width - sum(emu)
            hoehen = [int((n * ZEILE + 2 * oben) * groesse * EMU_PER_PT) for n in zeilen]
            return groesse, emu, hoehen, seite, oben
        groesse -= 1


def _zeilenzahlen(zellen: list[list[str]], breiten: list[float], groesse: float, seite: float, typeface: str,
                  kopf: bool) -> list[int]:
    """Lines per row, the header measured in bold."""
    aus = []
    for r, zeile in enumerate(zellen):
        font = schrift(typeface, groesse, fett=kopf and r == 0)
        aus.append(max(zeilen_hart(t, font, (b - 2 * seite * groesse) * _SPIEL) for t, b in zip(zeile, breiten)))
    return aus


def diagramm(slide, feld: Feld, body: Body, f: Farben, typeface: str) -> None:
    """Series and slices in steps of the one colour, light grid, a legend only where it tells something."""
    daten = CategoryChartData()
    daten.categories = body.categories
    for serie in body.series:
        daten.add_series(serie.name or " ", serie.values)
    art = body.chart_type if body.chart_type in _CHART else "bar"
    chart = slide.shapes.add_chart(_CHART[art], feld.left, feld.top, feld.width, feld.height, daten).chart
    groesse = min(_DIAGRAMM_PT[1], max(_DIAGRAMM_PT[0], round((feld.size_pt or 16) * _DIAGRAMM_ANTEIL)))
    chart.font.size = Pt(groesse)
    chart.font.color.rgb = RGBColor.from_string(fb.NEBENTEXT)
    if typeface:
        chart.font.name = typeface
    chart.has_title = False
    plot = chart.plots[0]
    punkte = len(body.categories) * len(body.series)
    if art == "pie":
        _torte(chart, plot, len(body.categories), f)
    else:
        _reihen(chart, plot, art, len(body.series), f)
        if punkte <= _MAX_BESCHRIFTETE_PUNKTE:
            plot.has_data_labels = True
            beschriftung = plot.data_labels
            beschriftung.number_format = "General"
            beschriftung.number_format_is_linked = False
            beschriftung.position = XL_LABEL_POSITION.OUTSIDE_END if art == "bar" else XL_LABEL_POSITION.ABOVE
            beschriftung.font.color.rgb = RGBColor(0, 0, 0)
            beschriftung.font.size = Pt(groesse)


def _reihen(chart, plot, art: str, n: int, f: Farben) -> None:
    for serie, farbe in zip(plot.series, f.stufen(n)):
        if art == "bar":
            serie.format.fill.solid()
            serie.format.fill.fore_color.rgb = farbe
            serie.format.line.fill.background()
        else:
            serie.smooth = False
            serie.format.line.color.rgb = farbe
            serie.format.line.width = Pt(2.5)
            serie.marker.style = XL_MARKER_STYLE.CIRCLE
            serie.marker.size = 7
            serie.marker.format.fill.solid()
            serie.marker.format.fill.fore_color.rgb = farbe
            serie.marker.format.line.color.rgb = farbe
    if art == "bar":
        plot.gap_width = 60
        plot.overlap = 0
    chart.has_legend = n > 1
    if n > 1:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    grau = RGBColor.from_string(fb.HILFSLINIE)
    wert = chart.value_axis
    wert.has_major_gridlines = True
    wert.major_gridlines.format.line.color.rgb = grau
    wert.major_gridlines.format.line.width = Pt(0.75)
    wert.format.line.fill.background()
    chart.category_axis.format.line.color.rgb = grau


def _torte(chart, plot, n: int, f: Farben) -> None:
    """Slices from full colour down to the lightest step, split by white edges, named in the legend."""
    for punkt, farbe in zip(plot.series[0].points, f.stufen(n)):
        punkt.format.fill.solid()
        punkt.format.fill.fore_color.rgb = farbe
        punkt.format.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        punkt.format.line.width = Pt(1.5)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    plot.has_data_labels = True
    beschriftung = plot.data_labels
    beschriftung.number_format = "General"
    beschriftung.number_format_is_linked = False
    beschriftung.position = XL_LABEL_POSITION.OUTSIDE_END
    beschriftung.font.color.rgb = RGBColor(0, 0, 0)
