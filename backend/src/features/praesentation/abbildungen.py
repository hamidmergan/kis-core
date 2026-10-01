"""Steps, key figures, a pyramid of levels and a set of symbols, laid out for whatever count they hold."""

from __future__ import annotations

import re

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from src.features.praesentation import farben as fb
from src.features.praesentation import icons, raster
from src.features.praesentation.elemente import _SPIEL, MIN_PT, ZEILE, schrift, setze_lauf
from src.features.praesentation.farben import Farben
from src.features.praesentation import formen
from src.features.praesentation.formen import _SCHRITT, _SCHWARZ, _emu, _faerbe, _form, _textfeld
from src.features.praesentation.schemas import Entry, Figure, Symbol
from src.features.praesentation.vorlage import EMU_PER_PT, Feld

# A step's circle in multiples of the base size, at most this share of its column, its number inside.
_KREIS = 1.2
_KREIS_ANTEIL = 0.5
_NUMMER = 0.45
# The line joining the steps, as a share of the circle, and the air between steps below each other.
_LINIE = 0.08
_ZEILENLUFT = 0.6
# A key figure at most this many times the base size, the lines between them this thin.
_MAX_ZAHL = 3.6
_TRENNER_PT = 1.5
# A pyramid is about as tall as an equilateral one of its base, the lines to its texts this thin.
_PYRAMIDE_HOEHE = 0.87
_VERBINDER_PT = 1.0
# A symbol's circle in multiples of the base size, at most this share of its column, the icon this share of it.
_SYMBOLKREIS = 4.0
_SYMBOLKREIS_ANTEIL = 0.6
_ICON_IM_KREIS = 0.52
# A drawing sits a third of the free height below its field's top, closer to the heading than to the foot.
_OBEN_ANTEIL = 0.35
# A cited provision breaks as a whole, a section sign at least keeps its number.
_KLAMMER = re.compile(r"\((?:§|Art\.)[^()]*\)")
_VERWEIS = re.compile(r"(§§?|Art\.)\s+(?=\d)")
_GESCHUETZT = "\u00a0"
# Measured in place of the protected space, a letter as wide as a space that no word wrap splits at.
_ALS_WORT = "i"


def _masse(feld: Feld) -> tuple[float, float, float, float]:
    return feld.left / EMU_PER_PT, feld.top / EMU_PER_PT, feld.width / EMU_PER_PT, feld.height / EMU_PER_PT


def _gebunden(text: str) -> str:
    text = _KLAMMER.sub(lambda m: m.group(0).replace(" ", _GESCHUETZT), text)
    return _VERWEIS.sub(lambda m: m.group(1) + _GESCHUETZT, text)


def _zeilen(text: str, typeface: str, groesse: float, breite: float, fett: bool = False) -> int:
    return formen._zeilen(_gebunden(text).replace(_GESCHUETZT, _ALS_WORT), typeface, groesse, breite, fett)


def _text(slide, x: float, y: float, w: float, h: float, text: str, groesse: float, farbe, fett: bool,
          typeface: str, ausrichtung=PP_ALIGN.CENTER, anker=MSO_ANCHOR.TOP):
    tf = _textfeld(slide, x, y, w, h, anker, ausrichtung)
    lauf = tf.paragraphs[0].add_run()
    lauf.text = _gebunden(text)
    setze_lauf(lauf, groesse, farbe, fett, typeface)
    return tf


def _ganz(texte: list[str], font, breite: float) -> bool:
    """No word of these texts is wider than the line, so none gets broken."""
    return all(font.getlength(w) <= breite * _SPIEL
               for t in texte for w in _gebunden(t).replace(_GESCHUETZT, _ALS_WORT).split())


def _groesster(messen, g: float, anteil: float = raster.TEXT_ANTEIL) -> tuple[bool, float, tuple]:
    """Whether it fits, the scale and the measures: the largest scale that fits, else the smallest readable one.

    The scale starts where the running text is as large as the template's body text.
    """
    k = 1 / anteil
    while True:
        passt, mass = messen(k)
        if passt or anteil * g * (k - _SCHRITT) < MIN_PT:
            return passt, k, mass
        k -= _SCHRITT


def _besser(a: tuple[bool, float, tuple], b: tuple[bool, float, tuple] | None) -> bool:
    """A layout that fits beats one that does not, among those the larger text wins, a tie keeps the earlier one."""
    return b is None or (a[0], a[1]) > (b[0], b[1] + 1e-9)


def _eintrag_hoehe(e: Entry, typeface: str, titel_pt: float, text_pt: float, breite: float) -> float:
    h = _zeilen(e.title, typeface, titel_pt, breite, True) * titel_pt * ZEILE
    if e.text:
        h += 0.3 * text_pt + _zeilen(e.text, typeface, text_pt, breite) * text_pt * ZEILE
    return h


def _eintrag(slide, x: float, y: float, breite: float, e: Entry, titel_pt: float, text_pt: float, f: Farben,
             typeface: str, ausrichtung=PP_ALIGN.CENTER) -> None:
    """Title in the colour and bold, the sentence below it in black."""
    titel_h = _zeilen(e.title, typeface, titel_pt, breite, True) * titel_pt * ZEILE
    _text(slide, x, y, breite, titel_h, e.title, titel_pt, f.textfarbe(), True, typeface, ausrichtung)
    if e.text:
        y += titel_h + 0.3 * text_pt
        h = _zeilen(e.text, typeface, text_pt, breite) * text_pt * ZEILE
        _text(slide, x, y, breite, h, e.text, text_pt, _SCHWARZ, False, typeface, ausrichtung)


# Steps


def ablauf(slide, feld: Feld, schritte: list[Entry], f: Farben, typeface: str) -> None:
    """Numbered circles on a line in reading order, side by side, or one below the other when that keeps the text larger."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    links, oben, breite, hoehe = _masse(feld)
    n = len(schritte)
    spalte = breite / n
    innen = spalte - g

    def ganz(titel_pt: float, text_pt: float, w: float) -> bool:
        return (_ganz([s.title for s in schritte], schrift(typeface, titel_pt, True), w)
                and _ganz([s.text for s in schritte], schrift(typeface, text_pt), w))

    def quer(k: float):
        titel_pt, text_pt = raster.TITEL_ANTEIL * g * k, raster.TEXT_ANTEIL * g * k
        d = min(_KREIS * g * k, spalte * _KREIS_ANTEIL)
        h = d + 0.5 * titel_pt + max(_eintrag_hoehe(s, typeface, titel_pt, text_pt, innen) for s in schritte)
        return ganz(titel_pt, text_pt, innen) and h <= hoehe, (titel_pt, text_pt, d, h, [], 0.0)

    def hoch(k: float):
        titel_pt, text_pt = raster.TITEL_ANTEIL * g * k, raster.TEXT_ANTEIL * g * k
        d = _KREIS * g * k
        text_w = breite - d - g
        zeilen = [max(d, _eintrag_hoehe(s, typeface, titel_pt, text_pt, text_w)) for s in schritte]
        luft = _ZEILENLUFT * g * k
        h = sum(zeilen) + (n - 1) * luft
        return ganz(titel_pt, text_pt, text_w) and h <= hoehe, (titel_pt, text_pt, d, h, zeilen, luft)

    nebeneinander = _groesster(quer, g)
    untereinander = _groesster(hoch, g)
    senkrecht = n > 1 and _besser(untereinander, nebeneinander)
    titel_pt, text_pt, d, h, zeilen, luft = (untereinander if senkrecht else nebeneinander)[2]
    oben += max(0.0, (hoehe - h) * _OBEN_ANTEIL)
    dicke = max(1.5, d * _LINIE)
    if senkrecht:
        tops = [oben + sum(zeilen[:i]) + i * luft for i in range(n)]
        _form(slide, MSO_SHAPE.RECTANGLE, links + (d - dicke) / 2, tops[0] + d / 2, dicke, tops[-1] - tops[0],
              f.ton(fb.LINIE))
        kreise = [(links, y) for y in tops]
    else:
        mitten = [links + spalte * (i + 0.5) for i in range(n)]
        if n > 1:
            _form(slide, MSO_SHAPE.RECTANGLE, mitten[0], oben + (d - dicke) / 2, mitten[-1] - mitten[0], dicke,
                  f.ton(fb.LINIE))
        kreise = [(x - d / 2, oben) for x in mitten]
    for i, (s, (x, y)) in enumerate(zip(schritte, kreise)):
        _form(slide, MSO_SHAPE.OVAL, x, y, d, d, f.ton(fb.VOLL))
        tf = _text(slide, x, y, d, d, str(i + 1), d * _NUMMER, f.schrift_auf(), True, typeface,
                   anker=MSO_ANCHOR.MIDDLE)
        tf.word_wrap = False
        if senkrecht:
            text_w = breite - d - g
            block = _eintrag_hoehe(s, typeface, titel_pt, text_pt, text_w)
            _eintrag(slide, x + d + g, y + max(0.0, (d - block) / 2), text_w, s, titel_pt, text_pt, f, typeface,
                     PP_ALIGN.LEFT)
        else:
            _eintrag(slide, x + d / 2 - innen / 2, y + d + 0.5 * titel_pt, innen, s, titel_pt, text_pt, f, typeface)


# Key figures


def kennzahlen(slide, feld: Feld, zahlen: list[Figure], f: Farben, typeface: str) -> None:
    """Each figure large in the colour with its label below, thin lines between the columns."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    links, oben, breite, hoehe = _masse(feld)
    spalte = breite / len(zahlen)
    innen = spalte - 1.5 * g
    # One size for every figure, the largest at which the widest one stays on one line.
    zahl_pt = min(_MAX_ZAHL * g, hoehe * 0.4)
    while zahl_pt > MIN_PT and max(schrift(typeface, zahl_pt, True).getlength(z.value)
                                   for z in zahlen) > innen * _SPIEL:
        zahl_pt -= 1

    def messen(k: float):
        text_pt = g * k
        zeilen = max(_zeilen(z.label, typeface, text_pt, innen) for z in zahlen)
        h = zahl_pt * ZEILE + 0.4 * text_pt + zeilen * text_pt * ZEILE
        return _ganz([z.label for z in zahlen], schrift(typeface, text_pt), innen) and h <= hoehe, (text_pt, h)

    text_pt, h = _groesster(messen, g, 1.0)[2]
    oben += max(0.0, (hoehe - h) * _OBEN_ANTEIL)
    for i, z in enumerate(zahlen):
        x = links + i * spalte
        if i:
            _form(slide, MSO_SHAPE.RECTANGLE, x - _TRENNER_PT / 2, oben, _TRENNER_PT, h, f.ton(fb.LINIE))
        tf = _text(slide, x + (spalte - innen) / 2, oben, innen, zahl_pt * ZEILE, z.value, zahl_pt, f.textfarbe(),
                   True, typeface)
        tf.word_wrap = False
        y = oben + zahl_pt * ZEILE + 0.4 * text_pt
        _text(slide, x + (spalte - innen) / 2, y, innen, oben + h - y, z.label, text_pt, _SCHWARZ, False, typeface)


# Levels


def _vieleck(slide, punkte: list[tuple[float, float]], farbe):
    """A closed polygon from points on the slide, the same outline in PowerPoint and LibreOffice."""
    x0, y0 = punkte[0]
    bauer = slide.shapes.build_freeform(_emu(x0), _emu(y0), scale=1.0)
    bauer.add_line_segments([(_emu(x), _emu(y)) for x, y in punkte[1:]], close=True)
    return _faerbe(bauer.convert_to_shape(), farbe)


def stufen(slide, feld: Feld, stufen_: list[Entry], f: Farben, typeface: str) -> None:
    """A pyramid in tints of the colour, the top level in full colour, each level's title and sentence beside it."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    links, oben, breite, hoehe = _masse(feld)
    n = len(stufen_)
    abstand = raster.STUFE_ABSTAND * g
    basis = breite * raster.PYRAMIDE_ANTEIL
    text_x = links + basis + 1.5 * g
    text_w = links + breite - text_x

    def messen(k: float):
        titel_pt, text_pt = raster.TITEL_ANTEIL * g * k, raster.TEXT_ANTEIL * g * k
        ganz = (_ganz([e.title for e in stufen_], schrift(typeface, titel_pt, True), text_w)
                and _ganz([e.text for e in stufen_], schrift(typeface, text_pt), text_w))
        stufe = max(raster.STUFE_MIN_HOEHE * g * k,
                    max(_eintrag_hoehe(e, typeface, titel_pt, text_pt, text_w) for e in stufen_))
        return ganz and n * stufe + (n - 1) * abstand <= hoehe, (titel_pt, text_pt, stufe)

    titel_pt, text_pt, stufe = _groesster(messen, g)[2]
    # The levels grow until the pyramid has its natural height, as far as the field allows.
    stufe = max(stufe, min((hoehe - (n - 1) * abstand) / n, (basis * _PYRAMIDE_HOEHE - (n - 1) * abstand) / n))
    h = n * stufe + (n - 1) * abstand
    oben += max(0.0, (hoehe - h) * _OBEN_ANTEIL)
    spitze = links + basis / 2

    def halb(y: float) -> float:
        return (y - oben) / h * basis / 2

    for i, (e, farbe) in enumerate(zip(stufen_, f.stufen(n))):
        ya = oben + i * (stufe + abstand)
        yb = ya + stufe
        kopf = [(spitze, ya)] if i == 0 else [(spitze - halb(ya), ya), (spitze + halb(ya), ya)]
        _vieleck(slide, kopf + [(spitze + halb(yb), yb), (spitze - halb(yb), yb)], farbe)
        mitte = ya + stufe / 2
        rand = spitze + halb(mitte) + 0.4 * g
        _form(slide, MSO_SHAPE.RECTANGLE, rand, mitte - _VERBINDER_PT / 2, text_x - 0.4 * g - rand, _VERBINDER_PT,
              f.ton(fb.LINIE))
        block = _eintrag_hoehe(e, typeface, titel_pt, text_pt, text_w)
        _eintrag(slide, text_x, mitte - block / 2, text_w, e, titel_pt, text_pt, f, typeface, PP_ALIGN.LEFT)


# Symbols


def symbole(slide, feld: Feld, symbole_: list[Symbol], f: Farben, typeface: str) -> None:
    """Icons in light circles with their term below, in the rows that keep the terms largest, each row centred."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    links, oben, breite, hoehe = _masse(feld)
    abstand = 1.5 * raster.ABSTAND * g

    def messen(reihen: list[int]):
        spalte = breite / max(reihen)
        innen = spalte - g

        def fuer(k: float):
            text_pt = g * k
            d = min(_SYMBOLKREIS * g * k, spalte * _SYMBOLKREIS_ANTEIL)
            zeilen = max(_zeilen(s.label, typeface, text_pt, innen, True) for s in symbole_)
            reihe = d + 0.5 * text_pt + zeilen * text_pt * ZEILE
            h = len(reihen) * reihe + (len(reihen) - 1) * abstand
            ganz = _ganz([s.label for s in symbole_], schrift(typeface, text_pt, True), innen)
            return ganz and h <= hoehe, (reihen, spalte, innen, text_pt, d, reihe, h)

        return fuer

    # The widest grid comes first, one with more rows wins only with larger terms.
    beste = None
    spalten_max = max(1, int(breite // (raster.SYMBOL_MIN_BREITE * g)))
    for spalten in range(min(len(symbole_), spalten_max), 0, -1):
        versuch = _groesster(messen(raster.verteilung(len(symbole_), spalten)), g, 1.0)
        if _besser(versuch, beste):
            beste = versuch
    reihen, spalte, innen, text_pt, d, reihe, h = beste[2]
    oben += max(0.0, (hoehe - h) * _OBEN_ANTEIL)
    farbe_hex = str(f.ton(fb.VOLL))
    icon = d * _ICON_IM_KREIS
    k = 0
    for r, anzahl in enumerate(reihen):
        start = links + (breite - anzahl * spalte) / 2
        y = oben + r * (reihe + abstand)
        for j, s in enumerate(symbole_[k:k + anzahl]):
            mitte = start + spalte * (j + 0.5)
            _form(slide, MSO_SHAPE.OVAL, mitte - d / 2, y, d, d, f.ton(fb.FLAECHE))
            slide.shapes.add_picture(icons.bild(s.icon, farbe_hex), _emu(mitte - icon / 2), _emu(y + (d - icon) / 2),
                                     _emu(icon), _emu(icon))
            unten = y + d + 0.5 * text_pt
            _text(slide, mitte - innen / 2, unten, innen, y + reihe - unten, s.label, text_pt, _SCHWARZ, True,
                  typeface)
        k += anzahl
