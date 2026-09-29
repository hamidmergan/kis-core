"""Cards, bars and the highlight box, drawn from shapes and laid out for whatever count they hold."""

from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu

from src.features.praesentation import farben as fb
from src.features.praesentation import icons, raster
from src.features.praesentation.elemente import _SPIEL, MIN_PT, ZEILE, schrift, setze_lauf, zeilen_hart
from src.features.praesentation.farben import Farben
from src.features.praesentation.schemas import Bar, Body, Card
from src.features.praesentation.vorlage import EMU_PER_PT, Feld

# Corner radius of a card as a share of its shorter side, as in the example deck.
_ECKE = 0.06
# A card's height against its width when the field has room, as in the example deck.
_MIN_SEITEN = 0.5
# Icon edge in multiples of the title size, and the paddings in multiples of the base size.
_ICON = 1.7
_RAND = 1.2
_SCHRITT = 0.05
# A bar's label column never takes more than this share of the field.
_MAX_BESCHRIFTUNG = 0.4
# Bar rows at their roomiest: pitch and bar height in multiples of the label size.
_MAX_ZEILENABSTAND = 3.2
_BALKEN_ANTEIL = 0.62
# The highlight box beside its sentence, and the size of that sentence against the body text.
_BOX_BREITE = 0.3
_BOX_SEITEN = 0.86
_AUSSAGE_ANTEIL = 1.9
_MAX_AUSSAGE_PT = 36.0
_SCHWARZ = RGBColor(0, 0, 0)
# Whole numbers in this range are read as years, the unit deadlines and durations arrive in.
_JAHRE = (1900, 2200)


def _emu(pt: float) -> Emu:
    return Emu(int(pt * EMU_PER_PT))


def _textfeld(slide, x: float, y: float, w: float, h: float, anker=MSO_ANCHOR.TOP, ausrichtung=PP_ALIGN.LEFT):
    """A bare text box in points: no margins, no autofit, so the measured layout is what renders."""
    kasten = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tf = kasten.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anker
    tf.paragraphs[0].alignment = ausrichtung
    return tf


def _form(slide, art, x: float, y: float, w: float, h: float, farbe, ecke: float | None = None):
    form = _faerbe(slide.shapes.add_shape(art, _emu(x), _emu(y), _emu(w), _emu(h)), farbe)
    if ecke is not None:
        form.adjustments[0] = ecke
    return form


def _faerbe(form, farbe):
    form.fill.solid()
    form.fill.fore_color.rgb = farbe
    form.line.fill.background()
    # The theme's shape style brings shadows LibreOffice draws despite an empty effect list, fill and line are set here.
    stil = form._element.find(qn("p:style"))
    if stil is not None:
        form._element.remove(stil)
    return form


def _zeilen(text: str, typeface: str, groesse: float, breite: float, fett: bool = False) -> int:
    return zeilen_hart(text, schrift(typeface, groesse, fett), breite * _SPIEL) if text else 0


# Cards


def karten(slide, feld: Feld, karten_: list[Card], f: Farben, typeface: str) -> None:
    """Rows of equal cards in the grid that keeps the text largest, the last row centred."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    breite = feld.width / EMU_PER_PT
    abstand = raster.ABSTAND * g
    reihen, karte_w, (titel_pt, text_pt, rand, icon, karte_h) = _kartenraster(karten_, feld, g, typeface)
    farbe_hex = str(f.ton(fb.VOLL))
    y = feld.top / EMU_PER_PT
    k = 0
    for anzahl in reihen:
        x = feld.left / EMU_PER_PT + (breite - anzahl * karte_w - (anzahl - 1) * abstand) / 2
        for karte in karten_[k:k + anzahl]:
            _form(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, karte_w, karte_h, f.ton(fb.FLAECHE), _ECKE)
            slide.shapes.add_picture(icons.bild(karte.icon, farbe_hex), _emu(x + rand), _emu(y + rand),
                                     _emu(icon), _emu(icon))
            innen = karte_w - 2 * rand
            oben = y + rand + icon + 0.5 * titel_pt
            titel_h = _zeilen(karte.title, typeface, titel_pt, innen, True) * titel_pt * ZEILE
            tf = _textfeld(slide, x + rand, oben, innen, titel_h)
            lauf = tf.paragraphs[0].add_run()
            lauf.text = karte.title
            setze_lauf(lauf, titel_pt, f.textfarbe(), True, typeface)
            if karte.text:
                oben += titel_h + 0.3 * text_pt
                tf = _textfeld(slide, x + rand, oben, innen, y + karte_h - rand - oben)
                lauf = tf.paragraphs[0].add_run()
                lauf.text = karte.text
                setze_lauf(lauf, text_pt, _SCHWARZ, False, typeface)
            x += karte_w + abstand
        k += anzahl
        y += karte_h + abstand


def _kartenraster(karten_: list[Card], feld: Feld, g: float, typeface: str) -> tuple[list[int], float, tuple]:
    """Cards per row, card width and measures, for the column count at which the text stays largest.

    The widest grid the field offers comes first, a grid with more columns wins only with larger
    text. When no grid fits, the one that overflows least is taken.
    """
    breite, hoehe = feld.width / EMU_PER_PT, feld.height / EMU_PER_PT
    abstand = raster.ABSTAND * g
    spalten_max, _ = raster.karten_raster(feld.width, feld.height, feld.size_pt)
    beste, notfall = None, None
    for spalten in range(min(spalten_max, len(karten_)), len(karten_) + 1):
        reihen = raster.verteilung(len(karten_), spalten)
        karte_w = (breite - abstand * (max(reihen) - 1)) / max(reihen)
        passt, faktor, mass = _kartenmass(karten_, karte_w, hoehe, len(reihen), abstand, g, typeface)
        kandidat = (reihen, karte_w, mass)
        if passt and (beste is None or faktor > beste[0] + 1e-9):
            beste = (faktor, kandidat)
        ueber = len(reihen) * mass[4] + (len(reihen) - 1) * abstand - hoehe
        if notfall is None or ueber < notfall[0] - 1e-9:
            notfall = (ueber, kandidat)
    return (beste or notfall)[1]


def _kartenmass(karten_: list[Card], karte_w: float, hoehe: float, reihen: int, abstand: float, g: float,
                typeface: str) -> tuple[bool, float, tuple[float, float, float, float, float]]:
    """Whether the rows fit, at which scale, and title size, text size, padding, icon edge and card height.

    The text starts at the template's body size and shrinks until every card fits without a broken
    word. A card is at least half as tall as it is wide when the field has the room, so two cards do
    not sit as thin strips above an empty slide.
    """
    platz = (hoehe - (reihen - 1) * abstand) / reihen
    faktor = 1 / raster.TEXT_ANTEIL
    while True:
        titel_pt = raster.TITEL_ANTEIL * g * faktor
        text_pt = raster.TEXT_ANTEIL * g * faktor
        rand = _RAND * g * faktor
        icon = _ICON * titel_pt
        innen = karte_w - 2 * rand
        titel_font, text_font = schrift(typeface, titel_pt, True), schrift(typeface, text_pt)
        ganz = all(titel_font.getlength(w) <= innen * _SPIEL for k in karten_ for w in k.title.split()) and all(
            text_font.getlength(w) <= innen * _SPIEL for k in karten_ for w in k.text.split())
        karte_h = 0.0
        for k in karten_:
            h = 2 * rand + icon + 0.5 * titel_pt + _zeilen(k.title, typeface, titel_pt, innen, True) * titel_pt * ZEILE
            if k.text:
                h += 0.3 * text_pt + _zeilen(k.text, typeface, text_pt, innen) * text_pt * ZEILE
            karte_h = max(karte_h, h)
        if ganz and karte_h <= platz:
            return True, faktor, (titel_pt, text_pt, rand, icon, max(karte_h, min(platz, karte_w * _MIN_SEITEN)))
        if text_pt - _SCHRITT * g < MIN_PT:
            return False, faktor, (titel_pt, text_pt, rand, icon, karte_h)
        faktor -= _SCHRITT


# Bars


def _zahl(wert: float) -> str:
    return f"{wert:g}"


def _basis(werte: list[float]) -> float:
    """Zero for amounts, for years the one before the earliest, so 2028 against 2030 does not read as 99 percent."""
    if werte and all(float(w).is_integer() and _JAHRE[0] <= w <= _JAHRE[1] for w in werte):
        return min(werte) - 1
    return 0.0


def balken(slide, feld: Feld, balken_: list[Bar], f: Farben, typeface: str) -> None:
    """Label left, a track in the light tint, the bar in full colour, its length measured on the largest value."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    breite, hoehe = feld.width / EMU_PER_PT, feld.height / EMU_PER_PT
    links, oben = feld.left / EMU_PER_PT, feld.top / EMU_PER_PT
    n = len(balken_)
    groesse = g
    while True:
        spalte = min(breite * _MAX_BESCHRIFTUNG,
                     max(schrift(typeface, groesse).getlength(b.label) for b in balken_) / _SPIEL + groesse)
        schritt = min(hoehe / n, _MAX_ZEILENABSTAND * groesse)
        zeilen = max(_zeilen(b.label, typeface, groesse, spalte - groesse) for b in balken_)
        if zeilen * groesse * ZEILE <= schritt * 0.95 or groesse - 1 < MIN_PT:
            break
        groesse -= 1
    dicke = min(schritt * _BALKEN_ANTEIL, 2 * groesse)
    spur_x = links + spalte
    spur_w = breite - spalte
    basis = _basis([b.value for b in balken_])
    groesstes = max((b.value for b in balken_), default=0.0) - basis
    wert_pt = max(MIN_PT, groesse * 0.85)
    for i, b in enumerate(balken_):
        mitte = oben + schritt * (i + 0.5)
        tf = _textfeld(slide, links, mitte - schritt / 2, spalte - groesse, schritt, MSO_ANCHOR.MIDDLE)
        lauf = tf.paragraphs[0].add_run()
        lauf.text = b.label
        setze_lauf(lauf, groesse, _SCHWARZ, False, typeface)
        y = mitte - dicke / 2
        _form(slide, MSO_SHAPE.ROUNDED_RECTANGLE, spur_x, y, spur_w, dicke, f.ton(fb.LINIE), 0.5)
        anteil = (b.value - basis) / groesstes if groesstes > 0 else 0.0
        laenge = max(dicke, spur_w * max(0.0, min(1.0, anteil)))
        _form(slide, MSO_SHAPE.ROUNDED_RECTANGLE, spur_x, y, laenge, dicke, f.ton(fb.VOLL), 0.5)
        text = b.text or _zahl(b.value)
        breit = schrift(typeface, wert_pt, True).getlength(text) + dicke
        if breit <= laenge:
            tf = _textfeld(slide, spur_x, y, laenge, dicke, MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
            farbe = f.schrift_auf()
        else:
            tf = _textfeld(slide, spur_x + laenge + dicke / 3, y, max(breit, spur_w - laenge - dicke / 3), dicke,
                           MSO_ANCHOR.MIDDLE)
            farbe = f.textfarbe()
        lauf = tf.paragraphs[0].add_run()
        lauf.text = text
        setze_lauf(lauf, wert_pt, farbe, True, typeface)


# Highlight


def hervorhebung(slide, feld: Feld, body: Body, f: Farben, typeface: str) -> None:
    """A box in full colour with the key word, beside it the sentence with its one term in colour."""
    typeface = typeface or "Calibri"
    g = raster.grund_pt(feld.size_pt)
    breite, hoehe = feld.width / EMU_PER_PT, feld.height / EMU_PER_PT
    links, oben = feld.left / EMU_PER_PT, feld.top / EMU_PER_PT
    box_w = breite * _BOX_BREITE
    box_h = min(hoehe, box_w * _BOX_SEITEN)
    box_x = links + breite - box_w
    _form(slide, MSO_SHAPE.ROUNDED_RECTANGLE, box_x, oben, box_w, box_h, f.ton(fb.VOLL), _ECKE)
    key_pt = box_h * 0.6
    while key_pt > MIN_PT and (schrift(typeface, key_pt, True).getlength(body.key) > box_w * 0.85
                               or key_pt * ZEILE > box_h * 0.9):
        key_pt -= 1
    tf = _textfeld(slide, box_x, oben, box_w, box_h, MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER)
    tf.word_wrap = False
    lauf = tf.paragraphs[0].add_run()
    lauf.text = body.key
    setze_lauf(lauf, key_pt, f.schrift_auf(), True, typeface)

    text_w = breite - box_w - 1.5 * g
    groesse = min(_MAX_AUSSAGE_PT, _AUSSAGE_ANTEIL * g)
    while groesse > MIN_PT and _zeilen(body.text, typeface, groesse, text_w, True) * groesse * ZEILE > box_h:
        groesse -= 1
    tf = _textfeld(slide, links, oben, text_w, box_h, MSO_ANCHOR.MIDDLE)
    for teil, betont in _teile(body.text, body.emphasis):
        lauf = tf.paragraphs[0].add_run()
        lauf.text = teil
        setze_lauf(lauf, groesse, f.textfarbe() if betont else _SCHWARZ, betont, typeface)


def _teile(text: str, betont: str) -> list[tuple[str, bool]]:
    """The sentence split around the first place the term occurs, ignoring case."""
    stelle = text.lower().find(betont.lower()) if betont.strip() else -1
    if stelle < 0:
        return [(text, False)]
    ende = stelle + len(betont)
    return [(t, b) for t, b in ((text[:stelle], False), (text[stelle:ende], True), (text[ende:], False)) if t]
