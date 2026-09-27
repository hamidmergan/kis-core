from io import BytesIO

from loguru import logger
from openai import OpenAIError
from pptx import Presentation

from src.services.provider_slots import ProviderFull
from src.services.translation import TranslationService
from src.utils.pptx_text import (
    flush_diagram_parts,
    iter_linked_text_nodes,
    iter_shape_paragraphs,
)
from src.utils.text import _RunRef, collect_paragraph_runs, translate_refs


def _collect_shape_text(shape, refs: list[_RunRef]) -> None:
    for para in iter_shape_paragraphs(shape):
        if para.text and para.text.strip():
            collect_paragraph_runs(para, refs)


def _collect_template_text(shape, refs: list[_RunRef]) -> None:
    """Text from master and layouts, without the placeholders."""
    if shape.is_placeholder:
        return
    _collect_shape_text(shape, refs)


def _collect_linked_text(slide, refs: list[_RunRef], touched_parts: set) -> None:
    """Collects text from SmartArt and chart parts."""
    for node in iter_linked_text_nodes(slide, touched_parts):
        text = node.text or ""
        if text.strip():
            leading = text[: len(text) - len(text.lstrip())]
            trailing = text[len(text.rstrip()):]
            refs.append(_RunRef(node, "text", leading, text.strip(), trailing))


def translate_pptx(
    content: bytes, source_lang: str, target_lang: str, service: TranslationService
) -> bytes:
    """Translates a PPTX preserving the layout, hidden slides included."""
    try:
        presentation = Presentation(BytesIO(content))
        refs: list[_RunRef] = []
        touched_parts: set = set()

        for slide in presentation.slides:
            for shape in slide.shapes:
                _collect_shape_text(shape, refs)

            # SmartArt and charts live in their own parts of the file.
            _collect_linked_text(slide, refs, touched_parts)

            if slide.has_notes_slide:
                for shape in slide.notes_slide.shapes:
                    _collect_shape_text(shape, refs)

        # Slide masters + layouts
        for slide_master in presentation.slide_masters:
            for shape in slide_master.shapes:
                _collect_template_text(shape, refs)
            for layout in slide_master.slide_layouts:
                for shape in layout.shapes:
                    _collect_template_text(shape, refs)

        if refs:
            translate_refs(refs, source_lang, target_lang, service)
            flush_diagram_parts(touched_parts)

        output = BytesIO()
        presentation.save(output)
        output.seek(0)
        return output.getvalue()

    except ValueError:
        raise
    except (OpenAIError, ProviderFull):
        # A busy or unreachable provider is no broken file, the job waits and tries again.
        raise
    except Exception as e:
        logger.exception("Error processing PPTX file")
        raise ValueError(f"Error processing PPTX file: {e}")
