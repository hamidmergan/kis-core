"""Document chunker for KI Tutor – extracts text chunks from PDF, DOCX, PPTX, TXT, and MD files."""

from loguru import logger

# Chunk sizes per embedding provider:
#   ollama / nomic-embed-text: BERT 2048-token ctx, German ~2-3 chars/token
#     → 600 chars ≈ 240 tokens, safely within limit
#   kisski / e5-mistral-7b-instruct: 4096-token ctx, German ~3-4 chars/token
#     → 1000 chars ≈ 250-333 tokens, well within limit and matches original value
_CHUNK_PARAMS: dict[str, tuple[int, int]] = {
    "ollama":  (600,  80),
    "kisski": (1000, 100),
}
_CHUNK_PARAMS_DEFAULT = _CHUNK_PARAMS["ollama"]


def _get_chunk_params() -> tuple[int, int]:
    """Return (chunk_size, overlap) for the configured embedding provider."""
    import os
    provider = os.getenv("TUTOR_EMBEDDING_PROVIDER", "ollama")
    return _CHUNK_PARAMS.get(provider, _CHUNK_PARAMS_DEFAULT)


def _split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    """Split text into overlapping chunks of chunk_size chars, breaking at whitespace."""
    text = text.strip()
    if len(text) <= chunk_size:
        return [text] if text else []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        # Break at last whitespace to avoid cutting words
        if end < len(text):
            last_space = chunk.rfind(" ")
            if last_space > chunk_size // 2:
                chunk = chunk[:last_space]
        chunk = chunk.strip()
        if chunk:
            chunks.append(chunk)
        next_start = start + len(chunk) - overlap
        if next_start <= start:
            break
        start = next_start
    return chunks


def extract_chunks(file_path: str, file_type: str) -> list[dict]:
    """Extract text chunks from a document file.

    Args:
        file_path: Absolute path to the file on disk.
        file_type:  One of "pdf", "docx", "pptx", "txt", "md".

    Returns:
        List of dicts with keys:
          - page_number (int, 1-indexed)
          - content    (str)
    """
    chunk_size, overlap = _get_chunk_params()
    ft = file_type.lower().lstrip(".")
    if ft == "pdf":
        return _extract_pdf(file_path, chunk_size, overlap)
    if ft == "docx":
        return _extract_docx(file_path, chunk_size)
    if ft == "pptx":
        return _extract_pptx(file_path, chunk_size, overlap)
    if ft in ("txt", "md"):
        return _extract_text(file_path, chunk_size, overlap)
    raise ValueError(f"Unsupported file_type: {file_type!r}")


def get_page_count(file_path: str, file_type: str) -> int:
    """Return the number of logical pages / chunks in a document."""
    return len(extract_chunks(file_path, file_type))


# ── Format-specific extractors ────────────────────────────────────────────────

def _extract_pdf(file_path: str, chunk_size: int, overlap: int) -> list[dict]:
    import fitz  # PyMuPDF
    chunks: list[dict] = []
    doc = fitz.open(file_path)
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        for sub in _split_text(text, chunk_size, overlap):
            chunks.append({"page_number": page_num, "content": sub})
    doc.close()
    if not chunks:
        raise ValueError(
            "Das PDF enthält keinen lesbaren Text. "
            "Gescannte PDFs werden aktuell nicht unterstützt. "
            "Bitte konvertiere das PDF zuerst mit einem OCR-Tool (z.B. Adobe Acrobat oder online unter smallpdf.com)."
        )
    return chunks


def _extract_docx(file_path: str, chunk_size: int) -> list[dict]:
    from docx import Document
    doc = Document(file_path)
    chunks: list[dict] = []
    current = ""
    page_number = 1

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if len(current) + len(text) + 1 > chunk_size:
            chunks.append({"page_number": page_number, "content": current or "[Leere Seite]"})
            page_number += 1
            current = text
        else:
            current = (current + " " + text).strip() if current else text

    # flush last chunk
    if current or not chunks:
        chunks.append({"page_number": page_number, "content": current or "[Leere Seite]"})

    return chunks


def _extract_pptx(file_path: str, chunk_size: int, overlap: int) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(file_path)
    chunks: list[dict] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        parts.append(t)
        content = " ".join(parts).strip()
        if not content:
            chunks.append({"page_number": slide_num, "content": "[Leere Seite]"})
        else:
            for sub in _split_text(content, chunk_size, overlap):
                chunks.append({"page_number": slide_num, "content": sub})
    return chunks


def _is_toc_page(blocks: list[dict]) -> bool:
    """Detect if a page is a table of contents by looking for dotted leader patterns."""
    import re
    toc_pattern = re.compile(r'\.{4,}\s*\d+\s*$')  # "........12"
    matches = sum(1 for b in blocks if toc_pattern.search(b["content"]))
    return matches >= 3  # 3+ lines with dotted leaders = TOC page



def extract_page_texts(file_path: str) -> list[dict]:
    """Extract the visible text of each page from a PDF.

    Returns one entry per page with meaningful text content.
    Pages that are TOC, nearly empty, or contain only formulas/images are skipped.

    Returns list of dicts with keys:
      - page_number (int, 1-indexed)
      - content     (str)
    """
    import fitz
    doc = fitz.open(file_path)
    pages: list[dict] = []

    for page_num, page in enumerate(doc, start=1):
        page_dict = page.get_text("dict")

        # Collect all text blocks
        page_blocks = []
        for block in page_dict["blocks"]:
            if block["type"] != 0:  # skip images
                continue
            lines_text = []
            for line in block["lines"]:
                spans_text = "".join(span["text"] for span in line["spans"])
                lines_text.append(spans_text.strip())
            content = " ".join(ln for ln in lines_text if ln).strip()
            if content:
                page_blocks.append({"content": content, "bbox": block["bbox"]})

        if not page_blocks:
            continue

        # Skip TOC pages
        if _is_toc_page(page_blocks):
            logger.info("extract_page_texts: skipping TOC page {}", page_num)
            continue

        # Join all text on the page
        full_text = " ".join(b["content"] for b in page_blocks).strip()

        # Strip leading page number (e.g. "11 Arbeitspaket..." → "Arbeitspaket...")
        import re
        full_text = re.sub(rf'^{page_num}\s+', '', full_text)
        word_count = len(full_text.split())

        # Skip pages with very little text (e.g. only formulas, images, page numbers)
        if word_count < 15:
            logger.info("extract_page_texts: skipping page {} ({} words)", page_num, word_count)
            continue

        pages.append({
            "page_number": page_num,
            "content": full_text,
        })

    doc.close()
    logger.info("extract_page_texts: {} pages with text extracted", len(pages))
    return pages


def _extract_text(file_path: str, chunk_size: int, overlap: int) -> list[dict]:
    with open(file_path, encoding="utf-8", errors="replace") as f:
        text = f.read()

    if not text.strip():
        return [{"page_number": 1, "content": "[Leere Seite]"}]

    chunks: list[dict] = []
    start = 0
    page_number = 1
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip() or "[Leere Seite]"
        chunks.append({"page_number": page_number, "content": chunk})
        page_number += 1
        start = end - overlap  # overlap with previous chunk
        if start >= len(text):
            break

    return chunks
