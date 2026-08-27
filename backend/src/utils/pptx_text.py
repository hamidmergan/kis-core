"""Shared text access to PPTX, SmartArt and charts included."""

from collections.abc import Iterator
from typing import Any

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE

_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_A_T = f"{{{_A_NS}}}t"
_C_V = f"{{{_C_NS}}}v"
_C_STR_CACHE = f"{{{_C_NS}}}strCache"

_DIAGRAM_DATA_RELTYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"
)

# The parsed tree is cached so reading and writing see the same thing.
_TREE_ATTR = "_kis_diagram_tree"


# ── Folien ──────────────────────────────────────────────────────────────────

def is_hidden_slide(slide) -> bool:
    """True when the slide is hidden in PowerPoint (show="0")."""
    return slide._element.get("show") == "0"


def iter_visible_slides(presentation) -> Iterator[Any]:
    """Only the slides LibreOffice actually exports."""
    for slide in presentation.slides:
        if not is_hidden_slide(slide):
            yield slide


# ── Paragraphs in shapes ────────────────────────────────────────────────────

def iter_shape_paragraphs(shape) -> Iterator[Any]:
    """All paragraphs of a shape, table cells and nested groups included."""
    if shape.has_text_frame:
        yield from shape.text_frame.paragraphs
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                yield from cell.text_frame.paragraphs
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for subshape in shape.shapes:
            yield from iter_shape_paragraphs(subshape)


def iter_slide_paragraphs(slide) -> Iterator[Any]:
    """Paragraphs of all shapes of a slide."""
    for shape in slide.shapes:
        yield from iter_shape_paragraphs(shape)


# ── Text in verlinkten Teilen (SmartArt, Diagramme) ─────────────────────────

def _diagram_tree(part):
    tree = getattr(part, _TREE_ATTR, None)
    if tree is None:
        tree = etree.fromstring(part.blob)
        setattr(part, _TREE_ATTR, tree)
    return tree


def _iter_diagram_nodes(slide, touched_parts: set | None) -> Iterator[Any]:
    """<a:t> elements from the data part of every SmartArt diagram of the slide."""
    for rel in slide.part.rels.values():
        if rel.reltype != _DIAGRAM_DATA_RELTYPE or rel.is_external:
            continue
        part = rel.target_part
        try:
            tree = _diagram_tree(part)
        except etree.XMLSyntaxError:
            continue
        if touched_parts is not None:
            touched_parts.add(part)
        yield from tree.iter(_A_T)


def _iter_chart_nodes(shape) -> Iterator[Any]:
    """The visible text of a chart, numbers stay out."""
    chart_element = shape.chart.part._element
    yield from chart_element.iter(_A_T)
    for value in chart_element.iter(_C_V):
        cache = value.getparent().getparent() if value.getparent() is not None else None
        if cache is not None and cache.tag == _C_STR_CACHE:
            yield value


def iter_linked_text_nodes(slide, touched_parts: set | None = None) -> Iterator[Any]:
    """Text nodes that do not live in the slide XML: SmartArt and charts."""
    yield from _iter_diagram_nodes(slide, touched_parts)
    for shape in _iter_all_shapes(slide.shapes):
        if getattr(shape, "has_chart", False):
            yield from _iter_chart_nodes(shape)


def _iter_all_shapes(shapes) -> Iterator[Any]:
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_all_shapes(shape.shapes)


def flush_diagram_parts(touched_parts) -> None:
    """Writes changed SmartArt trees back into the file."""
    for part in touched_parts:
        tree = getattr(part, _TREE_ATTR, None)
        if tree is None:
            continue
        part._blob = etree.tostring(
            tree, xml_declaration=True, encoding="UTF-8", standalone=True
        )
