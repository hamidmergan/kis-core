"""Room below the heading, headings that run long, and bullets in fields whose layout shows none."""

from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.features.praesentation import bau, platz, vorlage
from src.features.praesentation.schemas import Bar, Body, Column, Slide
from src.features.praesentation.vorlage import Feld

THL = Path("src/features/praesentation/assets/thl.pptx").read_bytes()
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
KURZ = "Aufstellung der Forschung"
LANG = "Thematische Aufstellung der Forschung und Lehre am Institut für Interaktive Systeme im Zusammenspiel mit der Hochschule"
PUNKTE = ["Erster Punkt", "Zweiter Punkt", "Dritter Punkt"]


@pytest.fixture(scope="module")
def analyse():
    return vorlage.analysiere(THL)


def _folie(analyse, s: Slide):
    return Presentation(BytesIO(bau.baue(THL, analyse, [s], {}, "D", False))).slides[0]


def _ph(folie, idx):
    return next(ph for ph in folie.placeholders if ph.placeholder_format.idx == idx)


def _groesse(folie, kopf) -> float:
    """The heading's size on the slide, the template's own when the build left it alone."""
    lauf = _ph(folie, kopf.idx).text_frame.paragraphs[0].runs[0]
    return lauf.font.size.pt if lauf.font.size else kopf.size_pt


def _kopf(analyse, layout):
    return analyse.layout(layout).der_art("ueberschrift")[0]


def test_vorlage_liest_verankerung_und_festen_zeilenabstand(analyse):
    deckblatt = _kopf(analyse, 0)
    assert (deckblatt.anker, deckblatt.zeile_pt) == ("t", 44.0)
    assert deckblatt.inset_oben == deckblatt.inset_unten == 216000
    assert _kopf(analyse, 3).zeile_pt == 0.0


def test_zeilenhoehe_folgt_festem_abstand():
    fest = Feld(0, "ueberschrift", "TITLE", 0, 0, 100, 100, size_pt=44.0, zeile_pt=44.0)
    assert platz.zeilenhoehe(fest, 22.0) == 22.0
    assert platz.zeilenhoehe(Feld(0, "ueberschrift", "TITLE", 0, 0, 100, 100, size_pt=32.0), 20.0) == 24.0


def test_verschiebung_behaelt_den_unteren_rand_und_die_reihenfolge():
    kopf = Feld(0, "ueberschrift", "BODY", 0, 0, 1000, 100)
    zwischen = Feld(1, "zwischentitel", "BODY", 0, 110, 1000, 50)
    text = Feld(2, "fliesstext", "OBJECT", 0, 170, 1000, 830)
    neu = platz.verschoben([kopf, zwischen, text], kopf, unten=100, luft=100, folie_hoehe=2000)
    assert neu[1].top == 200 and neu[2].top > neu[1].top + neu[1].height - 1
    assert neu[2].top + neu[2].height == pytest.approx(1000, abs=2)
    assert platz.verschoben([kopf, zwischen, text], kopf, unten=0, luft=100, folie_hoehe=2000) == {}


def test_inhalt_beginnt_eine_zeile_unter_dem_titeltext(analyse):
    folie = _folie(analyse, Slide(id="s", role="text", layout=3, title=KURZ, body=Body(items=PUNKTE)))
    kopf = _kopf(analyse, 3)
    feld = analyse.layout(3).der_art("fliesstext")[0]
    unten = platz.kopf_unten(kopf, KURZ, "Calibri", _groesse(folie, kopf))
    assert _ph(folie, feld.idx).top >= unten + platz.luft([feld]) - 1
    assert _ph(folie, feld.idx).top + _ph(folie, feld.idx).height == feld.top + feld.height


@pytest.mark.parametrize("layout,rolle,body,spalten", [
    (3, "text", Body(items=PUNKTE), []),
    (3, "text", Body(kind="bars", bars=[Bar(label="A", value=1), Bar(label="B", value=2)]), []),
    (5, "zwei_spalten", None, [Column(heading="L", items=PUNKTE), Column(heading="R", items=PUNKTE)]),
])
def test_langer_titel_schiebt_den_inhalt_statt_ihn_zu_ueberdecken(analyse, layout, rolle, body, spalten):
    folie = _folie(analyse, Slide(id="s", role=rolle, layout=layout, title=LANG, body=body, columns=spalten))
    kopf = _kopf(analyse, layout)
    unten = platz.kopf_unten(kopf, LANG, "Calibri", _groesse(folie, kopf))
    inhalt = [sh for sh in folie.shapes if sh.top > kopf.top and sh.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER] + [
        ph for ph in folie.placeholders if ph.placeholder_format.idx != kopf.idx and ph.top > kopf.top
        and ph.placeholder_format.idx in {f.idx for f in platz.darunter(analyse.layout(layout).felder, kopf)}]
    assert inhalt and min(sh.top for sh in inhalt) >= unten


def test_deckblatt_untertitel_rueckt_nur_bei_ueberlauf(analyse):
    kurz = _folie(analyse, Slide(id="s", role="deckblatt", layout=0, title=KURZ, subtitle="Stand"))
    untertitel = analyse.layout(0).der_art("zwischentitel")[0]
    assert _ph(kurz, untertitel.idx).top == untertitel.top
    lang = _folie(analyse, Slide(id="s", role="deckblatt", layout=0, title=LANG, subtitle="Stand"))
    kopf = _kopf(analyse, 0)
    titel = _ph(lang, kopf.idx).text_frame.paragraphs[0]
    groesse = _groesse(lang, kopf)
    assert _ph(lang, untertitel.idx).top >= platz.kopf_unten(kopf, LANG, "Calibri", groesse) - 1
    if titel.runs[0].font.size:
        assert titel.line_spacing.pt == pytest.approx(platz.zeilenhoehe(kopf, groesse), abs=0.1)


def _zeichen(ph) -> list[str | None]:
    aus = []
    for absatz in ph.text_frame.paragraphs:
        ppr = absatz._p.pPr
        char = ppr.find(f"{_A}buChar") if ppr is not None else None
        aus.append(char.get("char") if char is not None else None)
    return aus


def test_liste_bekommt_zeichen_wo_das_layout_keine_zeigt(analyse):
    s = Slide(id="s", role="text_mit_bild", layout=4, title=KURZ,
              columns=[Column(items=PUNKTE), Column(items=["Nur ein Satz"])])
    folie = _folie(analyse, s)
    links, rechts = analyse.layout(4).der_art("fliesstext")
    assert _zeichen(_ph(folie, links.idx)) == ["•"] * 3
    assert _zeichen(_ph(folie, rechts.idx)) == [None]


def test_layout_mit_eigenen_zeichen_bleibt_unberuehrt(analyse):
    folie = _folie(analyse, Slide(id="s", role="text", layout=3, title=KURZ, body=Body(items=PUNKTE)))
    assert _zeichen(_ph(folie, analyse.layout(3).der_art("fliesstext")[0].idx)) == [None] * 3
