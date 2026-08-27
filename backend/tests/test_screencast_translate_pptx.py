"""What actually reaches the translator, SmartArt and charts included."""

from io import BytesIO

import pytest
from pptx import Presentation
from pptx.util import Inches

from tests._pptx_fixtures import (
    CHART_CATEGORIES,
    CHART_TITLE,
    SMARTART_TEXTS,
    chart_pptx,
    hidden_slide_pptx,
    multi_run_paragraph_pptx,
    smartart_pptx,
)

from src.services.screencast.pptx import extract_text_per_slide
from src.translators.pptx import translate_pptx


class FakeService:
    """Records what was to be translated and prefixes EN::."""

    def __init__(self):
        self.seen: list[str] = []

    def translate_batch(self, texts, source_lang, target_lang):
        self.seen.extend(texts)
        return [f"EN::{t}" for t in texts]


class CountingService(FakeService):
    """Also counts what really reaches the provider."""

    def __init__(self):
        super().__init__()
        self.calls_with_text: list[str] = []

    def translate_batch(self, texts, source_lang, target_lang):
        self.calls_with_text.extend(texts)
        return super().translate_batch(texts, source_lang, target_lang)


def _all_text(pptx_bytes: bytes) -> str:
    return "\n".join(extract_text_per_slide(pptx_bytes))


def test_translates_smartart_text():
    service = FakeService()
    out = translate_pptx(smartart_pptx(), "DE", "EN-US", service)

    for original in SMARTART_TEXTS:
        assert original in service.seen, f"nie zum Übersetzer geschickt: {original!r}"
    text = _all_text(out)
    for original in SMARTART_TEXTS:
        assert f"EN::{original}" in text


def test_translates_chart_title_and_categories():
    service = FakeService()
    out = translate_pptx(chart_pptx(), "DE", "EN-US", service)

    assert CHART_TITLE in service.seen
    for category in CHART_CATEGORIES:
        assert category in service.seen
    text = _all_text(out)
    assert f"EN::{CHART_TITLE}" in text


def test_chart_numbers_stay_untouched():
    """Numbers also sit in <c:v> but must not go into the translation."""
    service = FakeService()
    translate_pptx(chart_pptx(), "DE", "EN-US", service)
    assert "19.2" not in service.seen
    assert "21.4" not in service.seen


def test_hidden_slides_are_translated_too():
    """Hidden slides do not make it into the video but belong in the PPTX download."""
    service = FakeService()
    translate_pptx(hidden_slide_pptx(), "DE", "EN-US", service)
    assert "Folie B versteckt" in service.seen


def test_plain_text_still_works():
    service = FakeService()
    out = translate_pptx(smartart_pptx("Guten Morgen"), "DE", "EN-US", service)
    assert "EN::Guten Morgen" in _all_text(out)


def test_output_is_a_valid_pptx():
    service = FakeService()
    out = translate_pptx(smartart_pptx(), "DE", "EN-US", service)
    prs = Presentation(BytesIO(out))
    assert len(prs.slides) == 2


def test_master_and_layout_prompt_text_is_not_translated():
    """Placeholder text from master and layouts only costs quota."""
    service = FakeService()
    translate_pptx(smartart_pptx(), "DE", "EN-US", service)

    assert not [s for s in service.seen if "Click to edit" in s]
    assert "‹#›" not in service.seen
    assert "Second level" not in service.seen


def test_real_master_text_is_still_translated():
    """A real text box in the master is visible and has to be translated."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.2), Inches(6.8), Inches(4), Inches(0.4))
    box.text_frame.text = "Fachhochschule Musterstadt"
    # python-pptx cannot create a shape directly in the master, so it is moved.
    master = prs.slide_masters[0]
    master.shapes._spTree.append(box._element)
    buf = BytesIO()
    prs.save(buf)

    service = FakeService()
    translate_pptx(buf.getvalue(), "DE", "EN-US", service)
    assert "Fachhochschule Musterstadt" in service.seen


def test_identical_texts_are_sent_once():
    """Translating duplicate texts twice only costs quota."""
    prs = Presentation()
    for _ in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = "Gleicher Text"
    buf = BytesIO()
    prs.save(buf)

    service = CountingService()
    out = translate_pptx(buf.getvalue(), "DE", "EN-US", service)

    assert service.calls_with_text.count("Gleicher Text") == 1
    assert _all_text(out).count("EN::Gleicher Text") == 3


def test_multi_run_paragraph_is_translated_as_one_sentence():
    """PowerPoint splits sentences into runs, translated one by one it becomes gibberish."""
    service = FakeService()
    translate_pptx(multi_run_paragraph_pptx(), "DE", "EN-US", service)

    real = [s for s in service.seen if "Vortrag" in s or "beginnt" in s or "neun" in s]
    assert real == ["Der Vortrag beginnt um neun Uhr"]


def test_multi_run_result_lands_in_the_paragraph():
    service = FakeService()
    out = translate_pptx(multi_run_paragraph_pptx(), "DE", "EN-US", service)
    assert _all_text(out) == "EN::Der Vortrag beginnt um neun Uhr"


def test_short_translation_result_is_an_error():
    """If fewer translations come back than were sent, the rest stayed silently German."""

    class ShortService:
        def translate_batch(self, texts, source_lang, target_lang):
            return [f"EN::{t}" for t in texts[:-1]]

    with pytest.raises(ValueError):
        translate_pptx(smartart_pptx(), "DE", "EN-US", ShortService())
