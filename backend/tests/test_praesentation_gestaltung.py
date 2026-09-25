"""Tables, charts, cards, bars and the highlight: the template's colour, and a layout for any count."""

import copy
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from src.features.praesentation import bau, elemente, icons, raster, vorlage
from src.features.praesentation.schemas import Bar, Body, Card, Series, Slide

THL = Path("src/features/praesentation/assets/thl.pptx").read_bytes()
ROT = "E4003A"
WORT = "Verwaltung Forschung Lehre Studium Transfer Internationales Digitalisierung Nachhaltigkeit".split()


@pytest.fixture(scope="module")
def analyse():
    return vorlage.analysiere(THL)


@pytest.fixture(scope="module")
def schmal(analyse):
    """The THL analysis with every content field half as wide, the narrow case a two-column template brings."""
    a = copy.deepcopy(analyse)
    for l in a.layouts:
        for f in l.felder:
            if f.art == "fliesstext":
                f.width //= 2
    return a


def _folie(a, body):
    prs = Presentation(BytesIO(bau.baue(THL, a, [Slide(id="s", role="text", title="T", body=body)], {}, "D", False)))
    return prs.slides[0]


def _feld(a):
    """The text area of the content field, where every drawn element has to stay."""
    return bau.textflaeche(a.layout(3).der_art("fliesstext")[0])


def _formen(folie):
    return [sh for sh in folie.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]


def _fuellung(sh) -> str:
    return str(sh.fill.fore_color.rgb)


def _innerhalb(sh, feld, spiel=12700):
    return (sh.left >= feld.left - spiel and sh.top >= feld.top - spiel
            and sh.left + sh.width <= feld.left + feld.width + spiel
            and sh.top + sh.height <= feld.top + feld.height + spiel)


# Tables


def _tabelle(folie):
    return next(sh for sh in folie.shapes if sh.has_table)


def test_tabelle_kopf_in_voller_farbe_und_zeilen_im_wechsel(analyse):
    body = Body(kind="table", header=["Frist", "Datum"], rows=[["A", "1"], ["B", "2"], ["C", "3"]])
    t = _tabelle(_folie(analyse, body)).table
    assert str(t.cell(0, 0).fill.fore_color.rgb) == ROT
    assert str(t.cell(0, 0).text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert [str(t.cell(r, 0).fill.fore_color.rgb) for r in (1, 2, 3)] == ["FFFFFF", "FEF2F5", "FFFFFF"]
    assert t._tbl.tblPr.find(qn("a:tableStyleId")).text == elemente._OHNE_STIL


def test_tabelle_ohne_kopf_hat_keine_farbige_zeile(analyse):
    t = _tabelle(_folie(analyse, Body(kind="table", rows=[["A", "1"], ["B", "2"]]))).table
    assert str(t.cell(0, 0).fill.fore_color.rgb) == "FFFFFF"
    assert t.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold is False


@pytest.mark.parametrize("spalten,zeilen", [(2, 2), (4, 4), (8, 6)])
def test_tabelle_fuellt_die_breite_und_passt_in_die_hoehe(analyse, spalten, zeilen):
    body = Body(kind="table", header=[f"Kopf {WORT[c]}" for c in range(spalten)],
                rows=[[f"{WORT[(r + c) % 8]} {r}" for c in range(spalten)] for r in range(zeilen)])
    rahmen = _tabelle(_folie(analyse, body))
    feld = _feld(analyse)
    assert sum(c.width for c in rahmen.table.columns) == feld.width
    assert sum(r.height for r in rahmen.table.rows) <= feld.height


def test_tabelle_breite_folgt_dem_text(analyse):
    body = Body(kind="table", header=["Nr", "Beschreibung"],
                rows=[["1", "Eine deutlich längere Beschreibung dieser Zeile mit vielen Wörtern"]])
    spalten = _tabelle(_folie(analyse, body)).table.columns
    assert spalten[1].width > 2 * spalten[0].width


def test_tabelle_in_einer_zeile_teilt_die_breite_nach_textlaenge(analyse):
    body = Body(kind="table", header=["Plattform", "Lizenz", "Betrieb"],
                rows=[["Eigenentwicklung", "Open Source", "zentral"]])
    breiten = [c.width for c in _tabelle(_folie(analyse, body)).table.columns]
    assert breiten[0] > breiten[1] > breiten[2]


def test_tabelle_schrift_schrumpft_bis_kein_wort_bricht(schmal):
    body = Body(kind="table", header=[f"Spalte {w}" for w in WORT[:5]],
                rows=[[f"{WORT[(r + c) % 8]} {r}" for c in range(5)] for r in range(4)])
    t = _tabelle(_folie(schmal, body)).table
    groesse = t.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt
    font = elemente.schrift("Calibri", groesse, fett=True)
    for c, spalte in enumerate(t.columns):
        laengstes = max(font.getlength(w) for r in range(len(t.rows)) for w in t.cell(r, c).text.split())
        assert laengstes <= spalte.width / 12700 - groesse
    assert groesse < _feld(schmal).size_pt


def test_tabelle_hat_hoechstens_acht_spalten(analyse):
    body = Body(kind="table", rows=[[str(i) for i in range(12)]])
    assert len(_tabelle(_folie(analyse, body)).table.columns) == 8


# Charts


def _chart(folie):
    return next(sh for sh in folie.shapes if sh.has_chart).chart


def test_saeulen_in_stufen_und_legende_erst_ab_zwei_reihen(analyse):
    eine = _chart(_folie(analyse, Body(kind="chart", categories=["a", "b"], series=[Series(name="x", values=[1, 2])])))
    assert eine.has_legend is False
    assert str(eine.plots[0].series[0].format.fill.fore_color.rgb) == ROT
    body = Body(kind="chart", categories=["a", "b"], series=[Series(name=str(i), values=[1, 2]) for i in range(4)])
    vier = _chart(_folie(analyse, body))
    assert vier.has_legend is True
    farben = [str(s.format.fill.fore_color.rgb) for s in vier.plots[0].series]
    assert farben[0] == ROT and len(set(farben)) == 4


def test_linien_in_stufen(analyse):
    body = Body(kind="chart", chart_type="line", categories=["a", "b"],
                series=[Series(name="x", values=[1, 2]), Series(name="y", values=[2, 1])])
    reihen = _chart(_folie(analyse, body)).plots[0].series
    assert str(reihen[0].format.line.color.rgb) == ROT and str(reihen[1].format.line.color.rgb) != ROT


def test_torte_faerbt_jedes_stueck(analyse):
    body = Body(kind="chart", chart_type="pie", categories=list("abcde"), series=[Series(name="x", values=[1] * 5)])
    chart = _chart(_folie(analyse, body))
    stuecke = [str(p.format.fill.fore_color.rgb) for p in chart.plots[0].series[0].points]
    assert stuecke[0] == ROT and len(set(stuecke)) == 5 and chart.has_legend


# Cards


@pytest.mark.parametrize("body", [
    Body(kind="table", header=["A", "B"], rows=[["1", "2"]]),
    Body(kind="bars", bars=[Bar(label="Ein langer Name", value=1), Bar(label="B", value=2)]),
    Body(kind="cards", cards=[Card(title="A", text="x"), Card(title="B", text="y")]),
    Body(kind="highlight", key="MO", text="Ein Satz."),
])
def test_links_und_rechts_bleibt_der_innenabstand_frei(analyse, body):
    """Drawn content starts where the field's own text starts and ends where that text wraps."""
    feld = _feld(analyse)
    formen = [sh for sh in _folie(analyse, body).shapes if sh.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER]
    assert min(sh.left for sh in formen) == feld.left
    assert max(sh.left + sh.width for sh in formen) <= feld.left + feld.width + 1
    assert feld.left > analyse.layout(3).der_art("fliesstext")[0].left


def test_raster_verteilung():
    assert [raster.verteilung(n, 3) for n in range(1, 9)] == [
        [1], [2], [3], [2, 2], [3, 2], [3, 3], [3, 2, 2], [3, 3, 2]]
    assert raster.verteilung(3, 1) == [1, 1, 1]


def test_platz_folgt_dem_feld(analyse, schmal):
    breit = vorlage.form_fuer(analyse.layout(3), analyse)
    eng = vorlage.form_fuer(schmal.layout(3), schmal)
    assert breit["karten"]["max"] == 6 and eng["karten"]["max"] < breit["karten"]["max"]
    assert breit["balken"] == raster.MAX_BALKEN


def _karten(n):
    return Body(kind="cards", cards=[Card(icon="lightbulb", title=f"Titel {i}", text=f"Ein Satz zur Karte {i}.")
                                     for i in range(n)])


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5, 6])
def test_karten_liegen_im_feld_gleich_gross_mit_icon(analyse, n):
    folie = _folie(analyse, _karten(n))
    karten = _formen(folie)
    bilder = [sh for sh in folie.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(karten) == n and len(bilder) == n
    assert len({(k.width, k.height) for k in karten}) == 1
    assert all(_innerhalb(k, _feld(analyse)) for k in karten)
    assert {_fuellung(k) for k in karten} == {"FCE6EB"}


def test_letzte_reihe_steht_mittig(analyse):
    karten = _formen(_folie(analyse, _karten(5)))
    feld = _feld(analyse)
    unten = [k for k in karten if k.top == max(k.top for k in karten)]
    assert len(unten) == 2
    links = unten[0].left - feld.left
    rechts = feld.left + feld.width - (unten[1].left + unten[1].width)
    assert abs(links - rechts) < 12700


@pytest.mark.parametrize("n", [1, 2])
def test_karten_passen_ins_schmale_feld(schmal, n):
    karten = _formen(_folie(schmal, _karten(n)))
    assert all(_innerhalb(k, _feld(schmal)) for k in karten)


def test_zu_viele_karten_bekommen_mehr_spalten_statt_herauszulaufen(schmal):
    karten = _formen(_folie(schmal, _karten(6)))
    assert len({k.left for k in karten}) > 1
    assert all(_innerhalb(k, _feld(schmal)) for k in karten)


def test_kartentitel_bricht_kein_wort(analyse):
    body = Body(kind="cards", cards=[Card(title="Forschungsdatenmanagement", text="x"),
                                     Card(title="B", text="y"), Card(title="C", text="z")])
    folie = _folie(analyse, body)
    titel = next(sh for sh in folie.shapes if sh.has_text_frame and sh.text_frame.text == "Forschungsdatenmanagement")
    groesse = titel.text_frame.paragraphs[0].runs[0].font.size.pt
    assert elemente.schrift("Calibri", groesse, fett=True).getlength("Forschungsdatenmanagement") <= titel.width / 12700


def test_icons_unbekannter_name_wird_ersatz():
    assert icons.gueltig("Light Bulb") == icons.ERSATZ
    assert icons.gueltig("LIGHTBULB") == "lightbulb"
    assert icons.gueltig("book_open") == "book-open"
    assert icons.gueltig(None) == icons.ERSATZ
    assert len(icons.namen()) >= 50


def test_icon_wird_eingefaerbt():
    from PIL import Image
    bild = Image.open(icons.bild("lightbulb", ROT))
    farben = {p[:3] for _, p in bild.getcolors(1 << 20) if p[3] == 255}
    assert farben == {(0xE4, 0x00, 0x3A)}


# Bars


def test_balken_laenge_folgt_dem_groessten_wert(analyse):
    body = Body(kind="bars", bars=[Bar(label="A", value=10), Bar(label="B", value=5), Bar(label="C", value=0)])
    formen = _formen(_folie(analyse, body))
    spuren = [f for f in formen if _fuellung(f) == "FACCD8"]
    balken = [f for f in formen if _fuellung(f) == ROT]
    assert len(spuren) == 3 and len(balken) == 3
    assert balken[0].width == spuren[0].width
    assert abs(balken[1].width - spuren[1].width / 2) < 12700
    assert balken[2].width == balken[2].height  # a zero still shows as a dot


def test_jahre_werden_ab_dem_jahr_vor_dem_fruehesten_gemessen(analyse):
    body = Body(kind="bars", bars=[Bar(label="A", value=2028), Bar(label="B", value=2029), Bar(label="C", value=2030)])
    formen = _formen(_folie(analyse, body))
    spur = next(f for f in formen if _fuellung(f) == "FACCD8").width
    laengen = [f.width for f in formen if _fuellung(f) == ROT]
    assert [round(l / spur, 2) for l in laengen] == [0.33, 0.67, 1.0]


def test_balkentext_steht_daneben_wenn_der_balken_zu_kurz_ist(analyse):
    body = Body(kind="bars", bars=[Bar(label="A", value=100, text="100 %"), Bar(label="B", value=1, text="1 %")])
    folie = _folie(analyse, body)
    texte = {sh.text_frame.text: sh for sh in folie.shapes if sh.has_text_frame}
    innen, aussen = texte["100 %"], texte["1 %"]
    assert str(innen.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert str(aussen.text_frame.paragraphs[0].runs[0].font.color.rgb) == ROT


@pytest.mark.parametrize("n", [2, 5, 8])
def test_balken_bleiben_im_feld(analyse, n):
    body = Body(kind="bars", bars=[Bar(label=f"Punkt {i}", value=i + 1) for i in range(n)])
    assert all(_innerhalb(f, _feld(analyse)) for f in _formen(_folie(analyse, body)))


def test_balken_ohne_text_zeigen_die_zahl(analyse):
    folie = _folie(analyse, Body(kind="bars", bars=[Bar(label="A", value=12.5), Bar(label="B", value=3)]))
    assert {"12.5", "3"} <= {sh.text_frame.text for sh in folie.shapes if sh.has_text_frame}


# Highlight


def test_hervorhebung_farbfeld_und_betontes_wort(analyse):
    body = Body(kind="highlight", key="MO", text="Ich bin montags vor Ort.", emphasis="Montags")
    folie = _folie(analyse, body)
    box = _formen(folie)[0]
    assert _fuellung(box) == ROT and _innerhalb(box, _feld(analyse))
    satz = next(sh for sh in folie.shapes if sh.has_text_frame and sh.text_frame.text == "Ich bin montags vor Ort.")
    laeufe = satz.text_frame.paragraphs[0].runs
    assert [l.text for l in laeufe] == ["Ich bin ", "montags", " vor Ort."]
    assert str(laeufe[1].font.color.rgb) == ROT and laeufe[1].font.bold
    schluessel = next(sh for sh in folie.shapes if sh.has_text_frame and sh.text_frame.text == "MO")
    assert str(schluessel.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"


def test_langes_schluesselwort_wird_kleiner(analyse):
    def groesse(key):
        folie = _folie(analyse, Body(kind="highlight", key=key, text="Satz."))
        sh = next(sh for sh in folie.shapes if sh.has_text_frame and sh.text_frame.text == key)
        return sh.text_frame.paragraphs[0].runs[0].font.size.pt

    assert groesse("Digitalisierung") < groesse("MO")
