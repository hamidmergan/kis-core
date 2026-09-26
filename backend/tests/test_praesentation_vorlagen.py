# backend/tests/test_praesentation_vorlagen.py
"""Own templates: what gets rejected, and that the system template always steps in."""

import json
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from pptx import Presentation

from src.features.praesentation import vorlagen

THL = Path("src/features/praesentation/assets/thl.pptx").read_bytes()


def _mit_makro(data: bytes) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for n in src.namelist():
            dst.writestr(n, src.read(n))
        dst.writestr("ppt/vbaProject.bin", b"x")
    return out.getvalue()


def _als_potx(data: bytes) -> bytes:
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for n in src.namelist():
            b = src.read(n)
            if n == "[Content_Types].xml":
                b = b.replace(b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
            dst.writestr(n, b)
    return out.getvalue()


class TestPruefung:
    def test_pptx_wird_angenommen(self):
        data, a = vorlagen.pruefe(THL, ".pptx")
        assert len(a.layouts) == 51 and data == THL

    def test_potx_wird_zu_pptx(self):
        data, a = vorlagen.pruefe(_als_potx(THL), ".potx")
        assert len(a.layouts) == 51
        assert len(Presentation(BytesIO(data)).slide_masters) == 1

    @pytest.mark.parametrize("ext", [".pptm", ".ppsx", ".pdf"])
    def test_falsche_endung(self, ext):
        with pytest.raises(vorlagen.VorlageUnbrauchbar) as e:
            vorlagen.pruefe(THL, ext)
        assert e.value.key == "praesentation.templateFormat"

    def test_makros_werden_abgewiesen(self):
        with pytest.raises(vorlagen.VorlageUnbrauchbar) as e:
            vorlagen.pruefe(_mit_makro(THL), ".pptx")
        assert e.value.key == "praesentation.templateMacros"

    def test_kaputte_datei(self):
        with pytest.raises(vorlagen.VorlageUnbrauchbar) as e:
            vorlagen.pruefe(b"kein zip", ".pptx")
        assert e.value.key == "praesentation.templateBroken"

    def test_zu_gross(self, monkeypatch):
        monkeypatch.setattr(vorlagen, "MAX_BYTES", 10)
        with pytest.raises(vorlagen.VorlageUnbrauchbar) as e:
            vorlagen.pruefe(THL, ".pptx")
        assert e.value.key == "praesentation.templateTooBig"


def test_systemvorlage_hat_stiltext():
    data, meta = vorlagen.systemvorlage()
    assert data[:2] == b"PK" and "vermilion" in meta["style_text"] and meta["name"] == "THL-Vorlage"


class TestSpeicher:
    def test_schluessel_liegen_unter_der_person(self):
        assert vorlagen.vorlage_key(7, "abc") == "users/7/praesentation/vorlagen/abc.pptx"
        assert vorlagen.vorschau_key(7, "abc") == "users/7/praesentation/vorlagen/abc.png"

    def test_aufloesen_nimmt_thl_wenn_die_eigene_fehlt(self, monkeypatch):
        monkeypatch.setattr(vorlagen.vorlagen_repo, "get", lambda uid, vid: {"id": vid, "name": "Meine", "stiltext": "S"})
        monkeypatch.setattr(vorlagen, "lade", lambda uid, vid: None)
        data, a, meta, fallback = vorlagen.aufloesen(7, "abc")
        assert fallback is True and meta["name"] == "THL-Vorlage" and len(a.layouts) == 51

    def test_aufloesen_nimmt_thl_wenn_die_eigene_kaputt_ist(self, monkeypatch):
        monkeypatch.setattr(vorlagen.vorlagen_repo, "get", lambda uid, vid: {"id": vid, "name": "Meine", "stiltext": "S"})
        monkeypatch.setattr(vorlagen, "lade", lambda uid, vid: b"kaputt")
        _, _, meta, fallback = vorlagen.aufloesen(7, "abc")
        assert fallback is True and meta["name"] == "THL-Vorlage"

    def test_aufloesen_eigene(self, monkeypatch):
        monkeypatch.setattr(vorlagen.vorlagen_repo, "get", lambda uid, vid: {"id": vid, "name": "Meine", "stiltext": "S"})
        monkeypatch.setattr(vorlagen, "lade", lambda uid, vid: THL)
        beruehrt = []
        monkeypatch.setattr(vorlagen.vorlagen_repo, "touch", lambda vid: beruehrt.append(vid))
        _, a, meta, fallback = vorlagen.aufloesen(7, "abc")
        assert fallback is False and meta == {"name": "Meine", "style_text": "S"} and beruehrt == ["abc"]

    def test_aufloesen_thl(self):
        _, a, meta, fallback = vorlagen.aufloesen(7, "thl")
        assert fallback is False and meta["name"] == "THL-Vorlage"

    def test_aufloesen_eigene_auch_wenn_touch_scheitert(self, monkeypatch):
        monkeypatch.setattr(vorlagen.vorlagen_repo, "get", lambda uid, vid: {"id": vid, "name": "Meine", "stiltext": "S"})
        monkeypatch.setattr(vorlagen, "lade", lambda uid, vid: THL)

        def _kaputtes_touch(vid):
            raise RuntimeError("db weg")

        monkeypatch.setattr(vorlagen.vorlagen_repo, "touch", _kaputtes_touch)
        _, a, meta, fallback = vorlagen.aufloesen(7, "abc")
        assert fallback is False and meta == {"name": "Meine", "style_text": "S"}


def test_meta_wird_neben_der_vorlage_abgelegt(monkeypatch):
    geschrieben = {}
    monkeypatch.setattr(vorlagen, "save_key_file", lambda key, data: geschrieben.__setitem__(key, data))
    vorlagen.speichere(7, "abc", b"x", None, {"roles": ["text"], "hasImages": False})
    assert json.loads(geschrieben["users/7/praesentation/vorlagen/abc.json"]) == {"roles": ["text"], "hasImages": False}


def test_vorlagenschluessel_liegen_unter_der_speicherwurzel(lokal):
    """Without S3, saving and loading a template round-trip through a path under STORAGE_ROOT."""
    vorlagen.speichere(7, "abc", b"pptx-inhalt", b"png-inhalt", {"roles": ["text"], "hasImages": False})
    assert vorlagen.lade(7, "abc") == b"pptx-inhalt"
    assert vorlagen.lade_vorschau(7, "abc") == b"png-inhalt"
    assert vorlagen.lade_meta(7, "abc") == {"roles": ["text"], "hasImages": False}
    assert (lokal / "users/7/praesentation/vorlagen/abc.pptx").read_bytes() == b"pptx-inhalt"
    assert (lokal / "users/7/praesentation/vorlagen/abc.png").read_bytes() == b"png-inhalt"


class _FakeS3Objekte:
    """A minimal S3 stand-in that keeps objects in memory, keyed like the real bucket."""

    def __init__(self):
        self.objekte: dict[str, bytes] = {}

    def put_object(self, Bucket=None, Key=None, Body=None, **kw):
        self.objekte[Key] = Body

    def get_object(self, Bucket=None, Key=None, **kw):
        daten = self.objekte[Key]

        class _Body:
            def read(self_inner):
                return daten
        return {"Body": _Body()}

    def head_object(self, Bucket=None, Key=None, **kw):
        from botocore.exceptions import ClientError
        if Key not in self.objekte:
            raise ClientError({"Error": {"Code": "404"}}, "HeadObject")

    def delete_object(self, Bucket=None, Key=None, **kw):
        self.objekte.pop(Key, None)


def test_vorlagenschluessel_gehen_auch_mit_s3(monkeypatch):
    """With S3 the same calls address the object store directly by key, no local path involved."""
    from src import storage
    client = _FakeS3Objekte()
    monkeypatch.setattr(storage, "USE_S3", True)
    monkeypatch.setattr(storage, "_s3", lambda: client)
    vorlagen.speichere(7, "abc", b"pptx-inhalt", b"png-inhalt", {"roles": ["text"], "hasImages": False})
    assert vorlagen.lade(7, "abc") == b"pptx-inhalt"
    assert vorlagen.lade_vorschau(7, "abc") == b"png-inhalt"
    assert vorlagen.lade_meta(7, "abc") == {"roles": ["text"], "hasImages": False}
    assert client.objekte["users/7/praesentation/vorlagen/abc.pptx"] == b"pptx-inhalt"
    vorlagen.loesche(7, "abc")
    assert vorlagen.lade(7, "abc") is None


def test_system_analyse_wird_nur_einmal_gerechnet(monkeypatch):
    vorlagen.system_analyse.cache_clear()
    original = vorlagen.vorlage.analysiere
    aufrufe = []
    monkeypatch.setattr(vorlagen.vorlage, "analysiere", lambda data: (aufrufe.append(1), original(data))[1])
    vorlagen.system_analyse()
    vorlagen.system_analyse()
    assert len(aufrufe) == 1


def test_deckblatt_vorschau_baut_eine_folie(monkeypatch):
    gesehen = []
    monkeypatch.setattr(vorlagen.vorschau, "render", lambda deck, dpi=80: gesehen.append(deck) or [b"png"])
    _, a = vorlagen.pruefe(THL, ".pptx")
    assert vorlagen.deckblatt_vorschau(THL, a, "Meine") == b"png"
    assert len(Presentation(BytesIO(gesehen[0])).slides) == 1


def test_system_cover_wird_nur_einmal_gerendert(monkeypatch):
    """The THL cover preview is expensive to render, one process must build it only once."""
    vorlagen.system_cover_vorschau.cache_clear()
    aufrufe = []
    monkeypatch.setattr(vorlagen.vorschau, "render", lambda deck, dpi=80: aufrufe.append(1) or [b"png"])
    vorlagen.system_cover_vorschau()
    vorlagen.system_cover_vorschau()
    assert len(aufrufe) == 1
