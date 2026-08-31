"""Text extraction from PPTX, with SmartArt, charts and hidden slides."""

from io import BytesIO

from pptx import Presentation
from pptx.oxml.ns import qn

from tests._pptx_fixtures import (
    CHART_CATEGORIES,
    CHART_TITLE,
    SMARTART_TEXTS,
    broken_linebreak_run_pptx,
    chart_pptx,
    hidden_slide_pptx,
    normautofit_textbox_pptx,
    smartart_pptx,
    spautofit_fits_pptx,
    spautofit_overflow_pptx,
    spautofit_textbox_pptx,
    tight_line_spacing_no_run_size_pptx,
    tight_line_spacing_percent_pptx,
    tight_line_spacing_pptx,
)

from src.services.screencast.pptx import (
    _apply_autofit_for_rendering,
    _fix_broken_linebreak_runs,
    _fix_tight_line_spacing,
    extract_text_per_slide,
)


def test_extracts_plain_textbox():
    texts = extract_text_per_slide(smartart_pptx())
    assert texts[0] == "Normaler Text auf Folie 1"


def test_extracts_smartart_text():
    texts = extract_text_per_slide(smartart_pptx())
    for expected in SMARTART_TEXTS:
        assert expected in texts[1], f"SmartArt-Text fehlt: {expected!r}"


def test_extracts_chart_title_and_categories():
    text = extract_text_per_slide(chart_pptx())[0]
    assert CHART_TITLE in text
    for category in CHART_CATEGORIES:
        assert category in text


def test_chart_numbers_are_not_extracted_as_text():
    """Only labels are text, the number series does not belong in the script."""
    text = extract_text_per_slide(chart_pptx())[0]
    assert "19.2" not in text
    assert "21.4" not in text


def test_hidden_slides_are_skipped():
    """LibreOffice does not render hidden slides, so they must not supply text either."""
    texts = extract_text_per_slide(hidden_slide_pptx())
    assert len(texts) == 2
    assert texts == ["Folie A", "Folie C"]


def test_smartart_text_appears_once_with_real_powerpoint_packaging():
    """The fallback group next to the chart would duplicate every line."""
    from tests._pptx_fixtures import smartart_with_fallback_pptx

    text = extract_text_per_slide(smartart_with_fallback_pptx())[1]
    assert text.count(SMARTART_TEXTS[0]) == 1
    for expected in SMARTART_TEXTS:
        assert expected in text


def test_fixes_broken_linebreak_run():
    """A bare '\\n' run with empty rPr becomes <a:br/> with borrowed formatting."""
    prs = Presentation(BytesIO(broken_linebreak_run_pptx()))

    changed = _fix_broken_linebreak_runs(prs)

    assert changed is True
    p = prs.slides[0].shapes[0].text_frame.paragraphs[0]._p
    runs = p.findall(qn("a:r"))
    breaks = p.findall(qn("a:br"))
    assert len(breaks) == 1
    assert all("\n" not in (r.find(qn("a:t")).text or "") for r in runs)
    br_rpr = breaks[0].find(qn("a:rPr"))
    assert br_rpr is not None
    assert br_rpr.get("sz") == "4350"


def test_fix_broken_linebreak_runs_is_noop_without_broken_runs():
    prs = Presentation(BytesIO(smartart_pptx()))

    changed = _fix_broken_linebreak_runs(prs)

    assert changed is False


def test_fixes_tight_line_spacing():
    """A <a:spcPts> smaller than the font's own size gets widened, not left to overlap."""
    prs = Presentation(BytesIO(tight_line_spacing_pptx()))

    changed = _fix_tight_line_spacing(prs)

    assert changed is True
    p_pr = prs.slides[0].shapes[0].text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    spc_pts = p_pr.find(qn("a:lnSpc")).find(qn("a:spcPts"))
    assert int(spc_pts.get("val")) == int(4000 * 1.2)


def test_fix_tight_line_spacing_is_noop_without_broken_paragraphs():
    prs = Presentation(BytesIO(smartart_pptx()))

    changed = _fix_tight_line_spacing(prs)

    assert changed is False


def test_fixes_tight_line_spacing_expressed_as_percentage():
    """A <a:spcPct> below single spacing (100%) is raised to the floor."""
    prs = Presentation(BytesIO(tight_line_spacing_percent_pptx()))

    changed = _fix_tight_line_spacing(prs)

    assert changed is True
    p_pr = prs.slides[0].shapes[0].text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    spc_pct = p_pr.find(qn("a:lnSpc")).find(qn("a:spcPct"))
    assert int(spc_pct.get("val")) == 100000


def test_fixes_tight_line_spacing_when_run_has_no_explicit_size():
    """A run without its own <a:rPr> size must still count towards the check,
    not be silently skipped."""
    prs = Presentation(BytesIO(tight_line_spacing_no_run_size_pptx()))

    changed = _fix_tight_line_spacing(prs)

    assert changed is True
    p_pr = prs.slides[0].shapes[0].text_frame.paragraphs[0]._p.find(qn("a:pPr"))
    spc_pts = p_pr.find(qn("a:lnSpc")).find(qn("a:spcPts"))
    assert int(spc_pts.get("val")) == int(1800 * 1.2)


def test_autofit_shrink_skips_spautofit_shapes(tmp_path):
    """spAutoFit shapes already grew to fit their text - shrinking the font
    would only shift where lines wrap versus the authored design."""
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(spautofit_textbox_pptx())

    render_path = _apply_autofit_for_rendering(pptx_path)

    rendered = Presentation(str(render_path))
    run = rendered.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 40


def test_autofit_shrink_applies_to_normautofit_shapes(tmp_path):
    """normAutofit shapes have a fixed box - LibreOffice needs the same fit
    correction PowerPoint would otherwise apply."""
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(normautofit_textbox_pptx())

    render_path = _apply_autofit_for_rendering(pptx_path)

    rendered = Presentation(str(render_path))
    run = rendered.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 32


def test_fixes_genuinely_overflowing_spautofit_text(tmp_path):
    """A spAutoFit box that measurably cannot fit its own text at native size
    gets shrunk just enough, unlike the blanket-guess shrink it replaced."""
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(spautofit_overflow_pptx())

    render_path = _apply_autofit_for_rendering(pptx_path)

    rendered = Presentation(str(render_path))
    run = rendered.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt < 18


def test_leaves_fitting_spautofit_text_at_native_size(tmp_path):
    """The same text in a tall-enough box must not be touched at all."""
    pptx_path = tmp_path / "deck.pptx"
    pptx_path.write_bytes(spautofit_fits_pptx())

    render_path = _apply_autofit_for_rendering(pptx_path)

    rendered = Presentation(str(render_path))
    run = rendered.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 18
