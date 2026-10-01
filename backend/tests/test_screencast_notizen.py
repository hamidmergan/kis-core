"""Speaker notes become scripts: one entry per visible slide, taken from the current deck."""

from io import BytesIO
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from src.services.screencast.pptx import extract_notes_per_slide, extract_text_per_slide
from src.storage import save_job_file

UUID = "d" * 32


def _deck(notizen: list[str | None], versteckt: set[int] = frozenset()) -> bytes:
    prs = Presentation()
    for i, notiz in enumerate(notizen):
        folie = prs.slides.add_slide(prs.slide_layouts[5])
        folie.shapes.title.text = f"Folie {i + 1}"
        if notiz is not None:
            folie.notes_slide.notes_text_frame.text = notiz
        if i in versteckt:
            folie._element.set("show", "0")
    puffer = BytesIO()
    prs.save(puffer)
    return puffer.getvalue()


def test_notizen_je_folie():
    deck = _deck(["Erste Notiz.\nZweite Zeile.", None, "Dritte."])
    assert extract_notes_per_slide(deck) == ["Erste Notiz.\nZweite Zeile.", "", "Dritte."]


def test_versteckte_folie_zaehlt_nicht_mit():
    deck = _deck(["Eins.", "Versteckt.", "Drei."], versteckt={1})
    notizen = extract_notes_per_slide(deck)
    assert notizen == ["Eins.", "Drei."]
    assert len(notizen) == len(extract_text_per_slide(deck))


@pytest.fixture
def client(lokal):
    from src.app import app
    from src.dependencies import get_current_user
    from src.dependencies_rbac import get_user_permissions
    from src.permissions import resolve_user_permissions
    from src.security.auth import verify_auth

    app.dependency_overrides[verify_auth] = lambda: None
    app.dependency_overrides[get_current_user] = lambda: {"user_id": 1, "email": "a@b.de", "role": "user"}
    app.dependency_overrides[get_user_permissions] = lambda: resolve_user_permissions("admin", None, True)
    jobs = {UUID: {"job_uuid": UUID, "user_id": 1}, "e" * 32: {"job_uuid": "e" * 32, "user_id": 2}}
    with patch("src.routes.screencast.scripts.db.get_job", jobs.get):
        yield TestClient(app)
    app.dependency_overrides.clear()


def test_route_liest_die_uebersetzte_datei(client):
    save_job_file(1, UUID, "translated.pptx", _deck(["Hallo.", None]))
    r = client.get(f"/screencast/{UUID}/notizen")
    assert r.status_code == 200
    assert r.json() == {"notes": ["Hallo.", ""]}


def test_korrigierte_datei_geht_vor(client):
    save_job_file(1, UUID, "translated.pptx", _deck(["Alt."]))
    save_job_file(1, UUID, "corrected.pptx", _deck(["Neu."]))
    assert client.get(f"/screencast/{UUID}/notizen").json() == {"notes": ["Neu."]}


def test_ohne_datei_404(client):
    assert client.get(f"/screencast/{UUID}/notizen").status_code == 404


def test_fremder_auftrag_404(client):
    assert client.get(f"/screencast/{'e' * 32}/notizen").status_code == 404


class _Uebersetzer:
    """Stands in for the translation provider and records what it was asked."""

    def __init__(self):
        self.aufrufe: list[tuple[list[str], str, str]] = []

    def translate_batch(self, texts, source_lang, target_lang):
        self.aufrufe.append((list(texts), source_lang, target_lang))
        return [f"{target_lang}: {t}" for t in texts]


def _mit_uebersetzer(uebersetzer):
    from src.services.glossary import GlossaryService
    return patch("src.dependencies.get_translate_client_for_provider",
                 lambda *a, **k: (uebersetzer, GlossaryService()))


def test_uebersprungene_uebersetzung_uebersetzt_die_notizen(client):
    from src.storage import save_result_meta
    save_job_file(1, UUID, "translated.pptx", _deck(["Hallo.", None, "Tschüss."]))
    save_result_meta(1, UUID, {"target_lang": "EN-US"})
    uebersetzer = _Uebersetzer()
    with _mit_uebersetzer(uebersetzer):
        r = client.get(f"/screencast/{UUID}/notizen")
    assert r.json() == {"notes": ["EN-US: Hallo.", "", "EN-US: Tschüss."]}
    assert uebersetzer.aufrufe == [(["Hallo.", "Tschüss."], "", "EN-US")]


def test_nach_der_uebersetzung_bleiben_die_notizen(client):
    from src.storage import save_result_meta
    save_job_file(1, UUID, "original.pptx", _deck(["Hallo."]))
    save_job_file(1, UUID, "translated.pptx", _deck(["Hello."]))
    save_result_meta(1, UUID, {"target_lang": "EN-US"})
    uebersetzer = _Uebersetzer()
    with _mit_uebersetzer(uebersetzer):
        assert client.get(f"/screencast/{UUID}/notizen").json() == {"notes": ["Hello."]}
    assert uebersetzer.aufrufe == []


def test_ohne_notizen_kein_aufruf(client):
    from src.storage import save_result_meta
    save_job_file(1, UUID, "translated.pptx", _deck([None, None]))
    save_result_meta(1, UUID, {"target_lang": "EN-US"})
    uebersetzer = _Uebersetzer()
    with _mit_uebersetzer(uebersetzer):
        assert client.get(f"/screencast/{UUID}/notizen").json() == {"notes": ["", ""]}
    assert uebersetzer.aufrufe == []


def test_gescheiterte_uebersetzung_meldet_sich(client):
    from src.storage import save_result_meta
    save_job_file(1, UUID, "translated.pptx", _deck(["Hallo."]))
    save_result_meta(1, UUID, {"target_lang": "EN-US"})
    with patch("src.dependencies.get_translate_client_for_provider", lambda *a, **k: (None, None)):
        r = client.get(f"/screencast/{UUID}/notizen")
    assert r.status_code == 502
    assert "übersetzen" in r.json()["detail"]
