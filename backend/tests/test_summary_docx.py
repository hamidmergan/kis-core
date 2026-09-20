"""The Word converter turns the summary's markdown into real Word structures."""

import io

from docx import Document

from src.features.summary import docx_export as de
from src.features.summary.docx_export import markdown_to_docx


def _open(data: bytes) -> Document:
    return Document(io.BytesIO(data))


class TestUeberschriften:
    def test_zwei_kreuz_wird_heading_2(self):
        doc = _open(markdown_to_docx("## Ein Titel"))
        assert doc.paragraphs[0].style.name == "Heading 2"
        assert doc.paragraphs[0].text == "Ein Titel"

    def test_alle_stufen_von_eins_bis_sechs(self):
        markdown = "\n".join(f"{'#' * n} Stufe {n}" for n in range(1, 7))
        doc = _open(markdown_to_docx(markdown))
        for absatz, n in zip(doc.paragraphs, range(1, 7)):
            assert absatz.style.name == f"Heading {n}"
            assert absatz.text == f"Stufe {n}"


class TestStichpunkte:
    def test_einfache_stichpunkte_mit_stern_und_strich(self):
        doc = _open(markdown_to_docx("* Erstens\n- Zweitens"))
        assert [a.style.name for a in doc.paragraphs] == ["List Bullet", "List Bullet"]
        assert [a.text for a in doc.paragraphs] == ["Erstens", "Zweitens"]

    def test_verschachtelte_stichpunkte_zwei_stufen(self):
        markdown = "- Oben\n    - Erste Stufe\n        - Zweite Stufe"
        doc = _open(markdown_to_docx(markdown))
        assert [a.style.name for a in doc.paragraphs] == [
            "List Bullet", "List Bullet 2", "List Bullet 3",
        ]
        assert [a.text for a in doc.paragraphs] == ["Oben", "Erste Stufe", "Zweite Stufe"]

    def test_nummerierte_liste_wird_list_number(self):
        doc = _open(markdown_to_docx("1. Erstens\n2. Zweitens"))
        assert [a.style.name for a in doc.paragraphs] == ["List Number", "List Number"]
        assert [a.text for a in doc.paragraphs] == ["Erstens", "Zweitens"]


class TestTabellen:
    def test_eine_tabelle_wird_eine_echte_word_tabelle(self):
        markdown = "| Name | Wert |\n| --- | --- |\n| A | 1 |\n| B | 2 |"
        doc = _open(markdown_to_docx(markdown))
        assert len(doc.tables) == 1
        table = doc.tables[0]
        assert len(table.rows) == 3
        assert [c.text for c in table.rows[0].cells] == ["Name", "Wert"]
        assert [c.text for c in table.rows[1].cells] == ["A", "1"]
        assert [c.text for c in table.rows[2].cells] == ["B", "2"]

    def test_die_kopfzeile_ist_fett(self):
        markdown = "| Name | Wert |\n| --- | --- |\n| A | 1 |"
        doc = _open(markdown_to_docx(markdown))
        header_cell = doc.tables[0].rows[0].cells[0]
        assert all(run.bold for run in header_cell.paragraphs[0].runs)
        data_cell = doc.tables[0].rows[1].cells[0]
        assert not any(run.bold for run in data_cell.paragraphs[0].runs)

    def test_zwei_tabellen_ohne_leerzeile_werden_zwei_tabellen(self):
        markdown = (
            "| A | B |\n| --- | --- |\n| 1 | 2 |\n"
            "| C | D |\n| --- | --- |\n| 3 | 4 |"
        )
        doc = _open(markdown_to_docx(markdown))
        assert len(doc.tables) == 2
        assert [c.text for c in doc.tables[0].rows[0].cells] == ["A", "B"]
        assert len(doc.tables[0].rows) == 2
        assert [c.text for c in doc.tables[1].rows[0].cells] == ["C", "D"]
        assert len(doc.tables[1].rows) == 2

    def test_zu_viele_spalten_werden_protokolliert(self, monkeypatch):
        meldungen = []
        monkeypatch.setattr(de.logger, "warning", lambda *a, **k: meldungen.append((a, k)))
        markdown = "| A | B |\n| --- | --- |\n| 1 | 2 | 3 |"
        doc = _open(markdown_to_docx(markdown))
        assert meldungen
        assert [c.text for c in doc.tables[0].rows[1].cells] == ["1", "2"]

    def test_zu_wenige_spalten_bleiben_leer(self):
        markdown = "| A | B |\n| --- | --- |\n| 1 |"
        doc = _open(markdown_to_docx(markdown))
        assert [c.text for c in doc.tables[0].rows[1].cells] == ["1", ""]


class TestFettschrift:
    def test_fett_mitten_im_satz(self):
        doc = _open(markdown_to_docx("Das ist **wichtig** und **noch etwas** dabei."))
        runs = doc.paragraphs[0].runs
        fett = [r.text for r in runs if r.bold]
        nicht_fett = [r.text for r in runs if not r.bold]
        assert fett == ["wichtig", "noch etwas"]
        assert "".join(runs_text for runs_text in [r.text for r in runs]) == (
            "Das ist wichtig und noch etwas dabei."
        )
        assert nicht_fett


class TestAbsaetze:
    def test_ein_gewoehnlicher_satz_wird_ein_absatz(self):
        doc = _open(markdown_to_docx("Ein ganz normaler Satz."))
        assert len(doc.paragraphs) == 1
        assert doc.paragraphs[0].text == "Ein ganz normaler Satz."
        assert doc.paragraphs[0].style.name in ("Normal", "Default Paragraph Style")

    def test_leerzeilen_erzeugen_keinen_leeren_absatz(self):
        doc = _open(markdown_to_docx("Erster Satz.\n\n\nZweiter Satz."))
        assert [a.text for a in doc.paragraphs] == ["Erster Satz.", "Zweiter Satz."]


class TestGemischtesDokument:
    def test_ueberschrift_liste_und_absatz_zusammen(self):
        markdown = "## Abschnitt\n\nEin Absatz davor.\n\n* Punkt eins\n* Punkt zwei\n\nEin Absatz danach."
        doc = _open(markdown_to_docx(markdown))
        texte = [(a.style.name, a.text) for a in doc.paragraphs]
        assert texte == [
            ("Heading 2", "Abschnitt"),
            ("Normal", "Ein Absatz davor."),
            ("List Bullet", "Punkt eins"),
            ("List Bullet", "Punkt zwei"),
            ("Normal", "Ein Absatz danach."),
        ]


class TestDeckblattUndVerzeichnis:
    """Mit Titel bekommt die Datei ein Deckblatt und ein Inhaltsverzeichnis."""

    def test_ohne_titel_beginnt_die_datei_beim_text(self):
        doc = _open(markdown_to_docx("Ein Satz."))
        assert doc.paragraphs[0].text == "Ein Satz."

    def test_der_titel_steht_auf_dem_deckblatt(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74", today="19.09.2026"))
        assert doc.paragraphs[0].text == "AMG74"
        assert doc.paragraphs[0].style.name == "Title"
        assert doc.paragraphs[1].text == "19.09.2026"

    def test_das_verzeichnis_ist_ein_feld(self):
        xml = _open(markdown_to_docx("## Kapitel", title="AMG74")).element.xml
        assert 'TOC \\o "1-3" \\h \\z \\u' in xml
        assert "fldChar" in xml

    def test_word_aktualisiert_beim_oeffnen_selbst(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74"))
        assert "updateFields" in doc.settings.element.xml

    def test_der_text_folgt_hinter_dem_vorspann(self):
        doc = _open(markdown_to_docx("## Kapitel\n\nEin Satz.", title="AMG74"))
        texte = [p.text for p in doc.paragraphs]
        assert "Kapitel" in texte
        assert texte.index("AMG74") < texte.index("Kapitel")


class TestHinweisUndVerzeichnisUeberschrift:
    """Deckblatt und Verzeichnis sprechen die Sprache des Dokuments, nicht die der Oberflaeche."""

    def test_der_hinweis_steht_direkt_unter_titel_und_datum(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74", today="19.09.2026"))
        assert doc.paragraphs[0].text == "AMG74"
        assert doc.paragraphs[1].text == "19.09.2026"
        assert doc.paragraphs[2].text.startswith("Mit künstlicher Intelligenz erstellt.")

    def test_der_hinweis_ist_kursiv(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74"))
        assert all(lauf.italic for lauf in doc.paragraphs[2].runs)

    def test_das_verzeichnis_hat_eine_ueberschrift(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74"))
        texte = [p.text for p in doc.paragraphs]
        assert "Inhaltsverzeichnis" in texte

    def test_die_ueberschrift_landet_nicht_im_verzeichnis(self):
        """Outline-Ebene 9 haelt sie aus dem TOC-Feld heraus, sonst listet es sich selbst."""
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74"))
        absatz = next(p for p in doc.paragraphs if p.text == "Inhaltsverzeichnis")
        assert absatz.style.name == "TOC Heading"
        assert 'w:val="9"' in absatz.style.element.xml

    def test_die_ueberschrift_steht_vor_dem_feld(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74"))
        texte = [p.text for p in doc.paragraphs]
        assert texte.index("Inhaltsverzeichnis") < texte.index("Kapitel")

    def test_englisch_bekommt_beide_texte_auf_englisch(self):
        doc = _open(markdown_to_docx("## Kapitel", title="AMG74", locale="en"))
        texte = [p.text for p in doc.paragraphs]
        assert "Table of contents" in texte
        assert doc.paragraphs[2].text.startswith("Created with artificial intelligence.")

    def test_ohne_titel_gibt_es_weder_hinweis_noch_verzeichnis(self):
        doc = _open(markdown_to_docx("Ein Satz."))
        assert [p.text for p in doc.paragraphs] == ["Ein Satz."]
