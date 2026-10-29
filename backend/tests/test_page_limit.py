"""The page limit applies equally to materials and tools."""

import pytest
from fastapi import HTTPException

from src.utils import file_validation


def _pptx(pfad, folien: int) -> str:
    """A real presentation with the requested number of empty slides."""
    from pptx import Presentation

    p = Presentation()
    for _ in range(folien):
        p.slides.add_slide(p.slide_layouts[6])
    ziel = str(pfad / "test.pptx")
    p.save(ziel)
    return ziel


def test_zaehlt_die_folien_einer_praesentation(tmp_path):
    assert file_validation.count_pages(".pptx", _pptx(tmp_path, 3)) == 3


@pytest.mark.parametrize("ext", [".txt", ".md", ".docx", ".mp3"])
def test_formate_ohne_seiten_liefern_nichts(tmp_path, ext):
    """Without a page count there is nothing to limit, that is not an error."""
    datei = tmp_path / f"x{ext}"
    datei.write_text("egal")
    assert file_validation.count_pages(ext, str(datei)) is None


def test_kaputte_datei_blockiert_den_upload_nicht(tmp_path):
    datei = tmp_path / "kaputt.pdf"
    datei.write_bytes(b"kein PDF")
    assert file_validation.count_pages(".pdf", str(datei)) is None


def _grenze(monkeypatch, wert: int) -> None:
    from src.services import tool_limits
    monkeypatch.setattr(tool_limits, "max_source_pages", lambda: wert)


def test_zu_viele_seiten_werden_abgelehnt(tmp_path, monkeypatch):
    _grenze(monkeypatch, 2)
    with pytest.raises(HTTPException) as fehler:
        file_validation.pages_or_raise(".pptx", _pptx(tmp_path, 3), "de")
    assert fehler.value.status_code == 413


def test_genau_die_grenze_geht_noch_durch(tmp_path, monkeypatch):
    _grenze(monkeypatch, 3)
    file_validation.pages_or_raise(".pptx", _pptx(tmp_path, 3), "de")


def test_ohne_seitenzahl_wird_nichts_abgelehnt(tmp_path, monkeypatch):
    _grenze(monkeypatch, 1)
    datei = tmp_path / "lang.txt"
    datei.write_text("x" * 100000)
    file_validation.pages_or_raise(".txt", str(datei), "de")
