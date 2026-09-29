"""The build fills only placeholders, places shapes where they belong and removes empty ones."""

from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from src.features.praesentation import bau, vorlage
from src.features.praesentation.schemas import Body, Column, ImageSpec, Series, Slide

THL = Path("src/features/praesentation/assets/thl.pptx").read_bytes()


@pytest.fixture(scope="module")
def analyse():
    return vorlage.analysiere(THL)


def _png(w=400, h=300, farbe=(200, 0, 0)) -> bytes:
    out = BytesIO()
    Image.new("RGB", (w, h), farbe).save(out, format="PNG")
    return out.getvalue()


def _deck(analyse, slides, bilder=None, notizen=False):
    return Presentation(BytesIO(bau.baue(THL, analyse, slides, bilder or {}, "Mein Deck", notizen)))


def _texte(slide) -> dict[int, str]:
    return {ph.placeholder_format.idx: ph.text_frame.text for ph in slide.placeholders if ph.has_text_frame}


def test_alte_folien_der_vorlage_sind_weg(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Deck", subtitle="U")])
    assert len(prs.slides) == 1


def test_deckblatt_titel_untertitel_und_banner(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Deck", subtitle="U",
                                images=[ImageSpec(prompt="x")])], {"s1_0": _png(1536, 400)})
    folie = prs.slides[0]
    assert folie.slide_layout.name == "Titelfolie 1"
    texte = _texte(folie)
    assert texte[0] == "Deck" and texte[1] == "U"
    assert 24 not in texte and 23 not in texte  # side texts are gone on the cover
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE
               and sh.image for sh in folie.shapes)


def test_textfolie_stichpunkte_und_randtext(analyse):
    prs = _deck(analyse, [Slide(id="s2", role="text", title="Kapitel", body=Body(items=["a", "b", "c"]))])
    folie = prs.slides[0]
    texte = _texte(folie)
    assert texte[13] == "Kapitel"
    assert texte[2].split("\n") == ["a", "b", "c"]
    assert texte[0] == "Mein Deck"  # the topmost side text carries the deck title
    assert 23 not in texte


def test_tabelle_sitzt_an_der_stelle_des_inhaltsfelds(analyse):
    body = Body(kind="table", header=["Jahr", "Wert"], rows=[["2020", "1"], ["2021", "2"]])
    prs = _deck(analyse, [Slide(id="s2", role="text", title="T", body=body)])
    folie = prs.slides[0]
    feld = bau.textflaeche(analyse.layout(3).der_art("fliesstext")[0])
    tabellen = [sh for sh in folie.shapes if sh.has_table]
    assert len(tabellen) == 1 and tabellen[0].left == feld.left and tabellen[0].top > feld.top
    assert tabellen[0].table.cell(0, 0).text == "Jahr" and tabellen[0].table.cell(2, 1).text == "2"
    assert 2 not in _texte(folie)


def test_diagramm_und_bild_im_inhaltsfeld(analyse):
    chart = Body(kind="chart", chart_type="pie", categories=["a", "b"], series=[Series(name="s", values=[1, 2])])
    bildkoerper = Body(kind="image", prompt="x")
    prs = _deck(analyse, [Slide(id="s2", role="text", title="T", body=chart),
                          Slide(id="s3", role="text", title="U", body=bildkoerper)], {"s3_body": _png()})
    assert any(sh.has_chart for sh in prs.slides[0].shapes)
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in prs.slides[1].shapes)


def test_stichpunkte_beginnen_unter_dem_titel_auch_im_eingerueckten_layout(analyse):
    schmal = next(l for l in analyse.layouts if l.name == "Text schmal")
    feld = schmal.der_art("fliesstext")[0]
    kopf = schmal.der_art("ueberschrift")[0]
    s = Slide(id="s2", role="text", layout=schmal.index, title="T", body=Body(items=["a", "b"]))
    ph = _deck(analyse, [s]).slides[0].placeholders[feld.idx]
    assert ph.left == kopf.left
    assert ph.left + ph.width == feld.left + feld.width


def test_diagramm_bleibt_im_eingerueckten_layout_mittig(analyse):
    schmal = next(l for l in analyse.layouts if l.name == "Text schmal")
    feld = bau.textflaeche(schmal.der_art("fliesstext")[0])
    chart = Body(kind="chart", categories=["a", "b"], series=[Series(name="s", values=[1, 2])])
    folie = _deck(analyse, [Slide(id="s2", role="text", layout=schmal.index, title="T", body=chart)]).slides[0]
    assert next(sh for sh in folie.shapes if sh.has_chart).left == feld.left


def test_zwei_spalten(analyse):
    s = Slide(id="s4", role="zwei_spalten", title="T",
              columns=[Column(heading="L", items=["1"]), Column(heading="R", items=["2", "3"])])
    texte = _texte(_deck(analyse, [s]).slides[0])
    assert (texte[1], texte[3]) == ("L", "R") and texte[2] == "1" and texte[4] == "2\n3"


def test_bild_wird_auf_das_feld_zugeschnitten(analyse):
    s = Slide(id="s5", role="text_mit_bild", layout=2, title="T", body=Body(items=["a"]),
              images=[ImageSpec(prompt="x"), ImageSpec(prompt="y")])
    prs = _deck(analyse, [s], {"s5_0": _png(1600, 400), "s5_1": _png(400, 1600)})
    bilder = [sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
              and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE and sh.image]
    assert len(bilder) == 2


def test_fehlendes_bild_laesst_kein_leeres_feld(analyse):
    s = Slide(id="s5", role="text_mit_bild", layout=2, title="T", body=Body(items=["a"]),
              images=[ImageSpec(prompt="x"), ImageSpec(prompt="y")])
    folie = _deck(analyse, [s]).slides[0]
    assert not any(sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE for sh in folie.placeholders)


def test_notizen_nur_wenn_gewuenscht(analyse):
    s = Slide(id="s2", role="text", title="T", body=Body(items=["a"]), notes="Sprich.")
    assert _deck(analyse, [s], notizen=True).slides[0].notes_slide.notes_text_frame.text == "Sprich."
    assert not _deck(analyse, [s], notizen=False).slides[0].has_notes_slide


def test_gespeichertes_layout_gewinnt_ueber_die_rotation(analyse):
    zaehler: dict[str, int] = {}
    assert bau.layout_index(analyse, Slide(id="s1", role="text_mit_bild", layout=4), zaehler) == 4
    assert analyse.layout(bau.layout_index(analyse, Slide(id="s2", role="text_mit_bild"), zaehler)).name == "Text und Bild"


def _hinweis_kaesten(slide) -> list:
    return [sh for sh in slide.shapes if sh.has_text_frame and not sh.is_placeholder]


def test_ki_hinweis_steht_auf_jeder_folie_unten_rechts(analyse):
    folien = [Slide(id="s1", role="deckblatt", title="Deck", subtitle="U"),
              Slide(id="s2", role="text", title="Inhalt", body=Body(kind="bullets", items=["A"])),
              Slide(id="s3", role="abschluss", title="Danke")]
    prs = Presentation(BytesIO(bau.baue(THL, analyse, folien, {}, "Mein Deck", False, hinweis="Mit KI erstellt")))
    for folie in prs.slides:
        kasten = _hinweis_kaesten(folie)
        assert len(kasten) == 1
        k = kasten[0]
        assert k.text_frame.text == "Mit KI erstellt"
        assert k.left + k.width <= prs.slide_width
        assert k.top > prs.slide_height * 0.85


def test_ki_hinweis_steht_auf_zehn_punkt_in_dunkelgrau(analyse):
    prs = Presentation(BytesIO(bau.baue(THL, analyse, [Slide(id="s1", role="text", title="T")], {},
                                        "Mein Deck", False, hinweis="Mit KI erstellt")))
    lauf = _hinweis_kaesten(prs.slides[0])[0].text_frame.paragraphs[0].runs[0]
    assert lauf.font.size.pt == bau.HINWEIS_PT == 10
    assert str(lauf.font.color.rgb) == "595959"


def test_ki_hinweis_bleibt_zehn_punkt_auch_ohne_randtext_in_der_vorlage():
    """The stock Office master has no edge text, the notice keeps its own size."""
    prs_leer = Presentation()
    roh = BytesIO()
    prs_leer.save(roh)
    daten = roh.getvalue()
    a = vorlage.analysiere(daten)
    prs = Presentation(BytesIO(bau.baue(daten, a, [Slide(id="s1", role="deckblatt", title="T")], {},
                                        "Mein Deck", False, hinweis="Mit KI erstellt")))
    lauf = _hinweis_kaesten(prs.slides[0])[0].text_frame.paragraphs[0].runs[0]
    assert lauf.font.size.pt == bau.HINWEIS_PT == 10


def test_ohne_hinweis_kommt_kein_kasten_dazu(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="text", title="T")])
    assert _hinweis_kaesten(prs.slides[0]) == []


def test_eingesetzte_bilder_landen_als_jpeg_in_der_datei(analyse):
    """The provider sends PNG, the deck must not carry it: eight of those push a file past 25 MB."""
    s = Slide(id="s5", role="text_mit_bild", layout=2, title="T", body=Body(items=["a"]),
              images=[ImageSpec(prompt="x"), ImageSpec(prompt="y")])
    prs = _deck(analyse, [s], {"s5_0": _png(1536, 1024), "s5_1": _png(1536, 1024)})
    bilder = [sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
              and sh.placeholder_format.type == PP_PLACEHOLDER.PICTURE and sh.image]
    assert len(bilder) == 2
    assert {sh.image.content_type for sh in bilder} == {"image/jpeg"}


def _groessen(folie, idx: int) -> list:
    ph = next(p for p in folie.placeholders if p.placeholder_format.idx == idx)
    return [lauf.font.size.pt if lauf.font.size else None for p in ph.text_frame.paragraphs for lauf in p.runs]


def test_kurzer_titel_behaelt_den_grad_der_vorlage(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Recht", subtitle="U")])
    assert _groessen(prs.slides[0], 0) == [None]


def test_langer_titel_schrumpft_statt_in_den_untertitel_zu_laufen(analyse):
    """In caps and bold this title is wider than its field, and the line below is the subtitle."""
    feld = analyse.layout(0).der_art("ueberschrift")[0]
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", subtitle="U",
                                title="Grundgesetz und wichtige Bundesgesetze")])
    (grad,) = _groessen(prs.slides[0], 0)
    assert grad is not None
    assert feld.size_pt * 0.6 <= grad < feld.size_pt


def test_der_grad_faellt_nie_unter_die_untergrenze(analyse):
    feld = analyse.layout(0).der_art("ueberschrift")[0]
    # The longest title the schema lets through at all.
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Verfassung " * 10, subtitle="U")])
    (grad,) = _groessen(prs.slides[0], 0)
    assert grad >= feld.size_pt * 0.6


def _nummer(folie) -> str | None:
    feld = [ph for ph in folie.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER]
    return feld[0].text_frame.text if feld else None


def test_foliennummer_unten_links_ausser_auf_dem_deckblatt(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Deck"),
                          Slide(id="s2", role="text", title="A", body=Body(items=["a"])),
                          Slide(id="s3", role="abschluss", title="Ende")])
    assert [_nummer(f) for f in prs.slides] == [None, "2", "3"]
    feld = next(ph for ph in prs.slides[1].placeholders
                if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER)
    assert feld.left < prs.slide_width // 4 and feld.top > prs.slide_height * 0.9
    # The field stays a slide number field, so PowerPoint renumbers after moving slides.
    assert 'type="slidenum"' in feld._element.xml


def test_ohne_datum_bleiben_datum_und_fusszeile_leer(analyse):
    prs = _deck(analyse, [Slide(id="s2", role="text", title="A", body=Body(items=["a"]))])
    arten = {ph.placeholder_format.type for ph in prs.slides[0].placeholders}
    assert PP_PLACEHOLDER.DATE not in arten and PP_PLACEHOLDER.FOOTER not in arten


def _datum(folie) -> str | None:
    feld = [ph for ph in folie.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.DATE]
    return feld[0].text_frame.text if feld else None


def test_datum_ueber_der_nummer_ausser_auf_dem_deckblatt(analyse):
    prs = Presentation(BytesIO(bau.baue(THL, analyse, [Slide(id="s1", role="deckblatt", title="Deck"),
                                                       Slide(id="s2", role="text", title="A", body=Body(items=["a"]))],
                                        {}, "Mein Deck", False, datum="24.09.2026")))
    assert [_datum(f) for f in prs.slides] == [None, "24.09.2026"]
    feld = next(ph for ph in prs.slides[1].placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.DATE)
    # A fixed text, a date field would jump to the day the file is opened.
    assert "datetime" not in feld._element.xml
    assert PP_PLACEHOLDER.FOOTER not in {ph.placeholder_format.type for ph in prs.slides[1].placeholders}


def test_datum_in_der_sprache_der_praesentation():
    from datetime import date
    assert bau.datum_text("de", date(2026, 9, 24)) == "24.09.2026"
    assert bau.datum_text("en", date(2026, 9, 24)) == "Sep 24, 2026"


def test_vorgegebene_nummern_fuer_eine_teilvorschau(analyse):
    prs = bau.baue(THL, analyse, [Slide(id="s5", role="text", title="A", body=Body(items=["a"]))], {},
                   "Mein Deck", False, nummern=[5])
    assert _nummer(Presentation(BytesIO(prs)).slides[0]) == "5"


def _punkt_groessen(folie, idx: int) -> set:
    return {r.font.size.pt if r.font.size else None
            for p in folie.placeholders[idx].text_frame.paragraphs for r in p.runs}


def test_passende_punkte_behalten_die_schrift_der_vorlage(analyse):
    prs = _deck(analyse, [Slide(id="s2", role="text", layout=3, title="A", body=Body(items=list("abcde")))])
    assert _punkt_groessen(prs.slides[0], 2) == {None}


def test_volle_punkte_schrumpfen_statt_ueber_den_rand_zu_laufen(analyse):
    lang = "Ein recht langer Stichpunkt mit vielen Woertern, der sicher zwei oder drei Zeilen braucht " * 4
    prs = _deck(analyse, [Slide(id="s2", role="text", layout=3, title="A", body=Body(items=[lang] * 8))])
    groessen = _punkt_groessen(prs.slides[0], 2)
    assert len(groessen) == 1 and None not in groessen
    assert min(groessen) >= 12


def test_volle_spalten_schrumpfen_auch(analyse):
    lang = "Ein recht langer Stichpunkt mit vielen Woertern, der sicher zwei Zeilen braucht"
    prs = _deck(analyse, [Slide(id="s4", role="text_mit_bild", layout=4, title="A",
                                columns=[Column(items=[lang] * 5), Column(items=["kurz"])])])
    assert None not in _punkt_groessen(prs.slides[0], 2)
    assert _punkt_groessen(prs.slides[0], 4) == {None}


def _fusszeile(folie) -> str | None:
    feld = [ph for ph in folie.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER]
    return feld[0].text_frame.text if feld else None


def test_name_unten_links_nur_auf_dem_deckblatt(analyse):
    prs = Presentation(BytesIO(bau.baue(THL, analyse, [Slide(id="s1", role="deckblatt", title="Deck"),
                                                       Slide(id="s2", role="text", title="A", body=Body(items=["a"]))],
                                        {}, "Mein Deck", False, vortragend="Prof. Dr. Erika Muster")))
    assert [_fusszeile(f) for f in prs.slides] == ["Prof. Dr. Erika Muster", None]
    feld = next(ph for ph in prs.slides[0].placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.FOOTER)
    assert feld.left < prs.slide_width // 4 and feld.top > prs.slide_height * 0.85


def test_ohne_namen_bleibt_das_deckblatt_unten_leer(analyse):
    prs = _deck(analyse, [Slide(id="s1", role="deckblatt", title="Deck")])
    assert _fusszeile(prs.slides[0]) is None
