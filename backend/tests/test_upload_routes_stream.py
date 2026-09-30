"""Every upload route streams its body to disk."""
import asyncio
from io import BytesIO

import pytest


class _ChunkOnlyUpload:
    """UploadFile stand-in whose ``read()`` refuses to hand over everything."""

    def __init__(self, data: bytes, filename: str):
        self.filename = filename
        self._data = data
        self._pos = 0

    async def read(self, n: int = -1) -> bytes:
        if n is None or n < 0:
            raise AssertionError("unbounded read() would buffer the whole body in RAM")
        chunk = self._data[self._pos:self._pos + n]
        self._pos += len(chunk)
        return chunk


class _Request:
    headers: dict = {}


def _docx_bytes(text: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.paragraphs[0].text = text
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _capture_storage(monkeypatch, module, attr="save_queue_file_from_path"):
    """Replace the path-based storage call and record what actually landed on disk."""
    stored: dict = {}

    def fake(user_id, job_uuid, ext_or_name, local_path):
        with open(local_path, "rb") as f:
            stored["bytes"] = f.read()
        return f"users/{user_id}/jobs/{job_uuid}/input"

    monkeypatch.setattr(module, attr, fake)
    return stored


def _silence_jobs(monkeypatch, module):
    monkeypatch.setattr(module.db, "create_job", lambda *a, **k: None)
    monkeypatch.setattr(module.db, "set_job_input_path", lambda *a, **k: None)


# --- per-route drivers -------------------------------------------------------

def _run_translate(monkeypatch, payload, filename):
    from src.routes import translate as route
    stored = _capture_storage(monkeypatch, route)
    _silence_jobs(monkeypatch, route)
    monkeypatch.setattr(route, "enforce_tool_burst", lambda *a, **k: None)
    monkeypatch.setattr(route, "enforce_quota", lambda *a, **k: None)
    monkeypatch.setattr(route, "is_admin_user", lambda u: False)

    import src.tasks.general_tasks as gt
    monkeypatch.setattr(gt, "run_translate", type("T", (), {"apply_async": staticmethod(
        lambda *a, **k: type("R", (), {"id": "task-1"})())})())
    monkeypatch.setattr(route.db, "set_job_celery_task_id", lambda *a, **k: None)

    asyncio.run(route.translate(
        request=_Request(), file=_ChunkOnlyUpload(payload, filename), text="",
        glossary="", provider_id=None, source_lang="", target_lang="",
        current_user={"user_id": 1}, locale="de",
    ))
    return stored



def _run_tts(monkeypatch, payload, filename):
    from src.routes import tts as route
    stored = _capture_storage(monkeypatch, route)
    _silence_jobs(monkeypatch, route)
    monkeypatch.setattr(route, "enforce_tool_burst", lambda *a, **k: None)
    monkeypatch.setattr(route, "enforce_quota", lambda *a, **k: None)
    monkeypatch.setattr(route, "is_admin_user", lambda u: False)
    monkeypatch.setattr(route.db, "set_job_celery_task_id", lambda *a, **k: None)
    # The voice settings for a later resume would otherwise land in the real storage.
    monkeypatch.setattr(route, "save_job_file", lambda *a, **k: "k")

    class _Client:
        @staticmethod
        def is_available():
            return True

    import src.dependencies as deps
    monkeypatch.setattr(deps, "get_tts_client_for_provider", lambda pid, user_id=None: _Client())
    monkeypatch.setattr(route.tts_job_provider, "is_there", lambda pid: True)

    import src.tasks.tts_tasks as tt
    monkeypatch.setattr(
        tt, "run_tts",
        type("T", (), {"apply_async": staticmethod(lambda *a, **k: type("R", (), {"id": "x"})())})(),
    )

    asyncio.run(route.text_to_speech(
        request=_Request(), file=_ChunkOnlyUpload(payload, filename), text="",
        voice_id="v", speed=1.0, provider_id=1, voice_label="",
        current_user={"user_id": 1}, locale="de",
    ))
    return stored


def _run_screencast(monkeypatch, payload, filename):
    from src.routes.screencast import translate as route
    stored = _capture_storage(monkeypatch, route)
    _silence_jobs(monkeypatch, route)
    monkeypatch.setattr(route, "validate_upload", lambda *a, **k: None)
    monkeypatch.setattr(route, "maybe_start_cleanup_task", lambda *a, **k: None)
    monkeypatch.setattr(route, "_enforce_queue_and_quota", lambda *a, **k: None)
    monkeypatch.setattr(route, "save_job_file", lambda *a, **k: "k")
    monkeypatch.setattr(route.db, "update_job_status", lambda *a, **k: None)

    import src.tasks.screencast_tasks as st
    monkeypatch.setattr(st, "dispatch_next_screencast_job_for_user", lambda uid: None)

    asyncio.run(route.screencast_translate(
        request=_Request(), file=_ChunkOnlyUpload(payload, filename),
        glossary_json="", source_lang="DE", target_lang="EN-US", provider_id=None,
        current_user={"user_id": 1}, locale="de",
    ))
    return stored


# Die Zusammenfassung fehlt hier bewusst: sie legt die Rohdatei nicht mehr ab,
# sie liest sie und speichert den Text. Siehe test_summary_route.py.
CASES = [
    ("translate/txt", _run_translate, lambda: b"Ein Satz zum Uebersetzen.\n" * 2000, "doc.txt"),
    ("translate/docx", _run_translate, lambda: _docx_bytes("Absatz"), "doc.docx"),
    ("tts/txt", _run_tts, lambda: b"Vorzulesender Text.\n" * 2000, "doc.txt"),
    ("tts/docx", _run_tts, lambda: _docx_bytes("Absatz"), "doc.docx"),
    ("screencast/pptx", _run_screencast, lambda: _pptx_bytes("Folientext"), "deck.pptx"),
]


@pytest.mark.parametrize("label,driver,make_payload,filename", CASES, ids=[c[0] for c in CASES])
def test_upload_reaches_storage_byte_for_byte_without_buffering(
    monkeypatch, label, driver, make_payload, filename
):
    payload = make_payload()
    stored = driver(monkeypatch, payload, filename)
    assert stored["bytes"] == payload
