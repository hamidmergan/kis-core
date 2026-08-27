"""The text extractors take a path and return the same text as from bytes."""
from io import BytesIO

import pytest
from docx import Document
from pptx import Presentation

from src.utils.text_extraction import EXTRACTORS


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes(lines: list[str]) -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    for i, line in enumerate(lines):
        para = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        para.text = line
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("ext,data,expected_fragment", [
    (".txt", b"hallo welt", "hallo welt"),
    (".md", b"# Titel\n\nText hier", "Titel"),
    (".vtt", b"WEBVTT\n\n00:00.000 --> 00:01.000\nUntertitel\n", "Untertitel"),
    (".srt", b"1\n00:00:00,000 --> 00:00:01,000\nZeile eins\n", "Zeile eins"),
])
def test_text_formats_read_from_path(tmp_path, ext, data, expected_fragment):
    path = tmp_path / f"file{ext}"
    path.write_bytes(data)

    from_path = EXTRACTORS[ext](str(path))
    from_bytes = EXTRACTORS[ext](data)

    assert expected_fragment in from_path
    assert from_path == from_bytes


def test_docx_reads_from_path(tmp_path):
    data = _docx_bytes(["Erste Zeile", "Zweite Zeile"])
    path = tmp_path / "doc.docx"
    path.write_bytes(data)

    from_path = EXTRACTORS[".docx"](str(path))

    assert "Erste Zeile" in from_path
    assert "Zweite Zeile" in from_path
    assert from_path == EXTRACTORS[".docx"](data)


def test_pptx_reads_from_path(tmp_path):
    data = _pptx_bytes(["Folientitel", "Folientext"])
    path = tmp_path / "deck.pptx"
    path.write_bytes(data)

    from_path = EXTRACTORS[".pptx"](str(path))

    assert "Folientitel" in from_path
    assert "Folientext" in from_path
    assert from_path == EXTRACTORS[".pptx"](data)
